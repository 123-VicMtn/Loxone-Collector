"""
Export du décompte sur une période choisie (`du` / `au`) : plusieurs mois,
ou des dates quelconques (départ d'un locataire le 15).

Site « Ref » (tests/site_reference.py) : septembre sur les séries du
Miniserver, octobre en répartition recalculée, App 2 sans relevé en octobre.
"""
from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

import site_reference as ref

URL = f"/api/decompte/export.{{}}?miniserver={ref.SITE}"


def _csv(client, query: str) -> list[list[str]]:
    res = client.get(URL.format("csv") + "&" + query)
    assert res.status_code == 200, res.get_data(as_text=True)
    return [ligne.split(";") for ligne in res.get_data().decode("utf-8-sig").splitlines()]


def _section(lignes: list[list[str]], titre: str) -> list[list[str]]:
    debut = lignes.index([titre]) + 1
    fin = next((i for i in range(debut, len(lignes)) if lignes[i] == [""]), len(lignes))
    return lignes[debut:fin]


def _kwh(texte: str) -> float | None:
    return float(texte.replace(",", ".")) if texte else None


def _par_zone(section: list[list[str]], colonne: int) -> dict[str, float | None]:
    return {row[1]: _kwh(row[colonne]) for row in section[1:]}


def test_un_mois_par_dates_identique_a_l_export_du_mois(client_user):
    assert _csv(client_user, "du=2026-10-01&au=2026-10-31") == _csv(client_user, "mois=2026-10")


def test_plusieurs_mois_additionnes(client_user):
    plage = _csv(client_user, "du=2026-09-01&au=2026-10-31")
    elec = _section(plage, "Électricité")
    assert elec[0][0] == "Période"
    assert elec[1][0] == "01.09.2026 → 31.10.2026"
    sept = _par_zone(_section(_csv(client_user, "mois=2026-09"), "Électricité"), 2)
    oct_ = _par_zone(_section(_csv(client_user, "mois=2026-10"), "Électricité"), 2)
    total = _par_zone(elec, 2)
    assert total["App 1"] == pytest.approx(sept["App 1"] + oct_["App 1"], abs=0.01)
    assert total["Communs"] == pytest.approx(sept["Communs"] + oct_["Communs"], abs=0.01)
    # Octobre sans relevé : jamais un total de période silencieusement incomplet.
    assert total["App 2"] is None
    assert elec[2][-1] == "données incomplètes"


def test_detail_par_mois_et_parts_de_chaque_mois(client_user):
    detail = _section(_csv(client_user, "du=2026-09-01&au=2026-10-31"), "Détail par mois")
    lignes = {(row[0], row[1]): row for row in detail[1:]}
    assert lignes[("Septembre 2026", "App 1")][5] == "Miniserver"
    assert lignes[("Octobre 2026", "App 1")][5] == "calculées"
    assert lignes[("Octobre 2026", "App 2")][6] == "données incomplètes"


def test_chauffage_lu_d_un_bout_a_l_autre(client_user):
    sept = _section(_csv(client_user, "mois=2026-09"), "Chauffage")
    oct_ = _section(_csv(client_user, "mois=2026-10"), "Chauffage")
    plage = _section(_csv(client_user, "du=2026-09-01&au=2026-10-31"), "Chauffage")
    assert plage[1][3] == sept[1][3]  # index de début : celui du 1er septembre
    assert plage[1][4] == oct_[1][4]  # index de fin : celui du 31 octobre


def test_depart_en_cours_de_mois(client_user):
    plage = _csv(client_user, "du=2026-09-15&au=2026-10-15")
    detail = _section(plage, "Détail par mois")
    assert {row[0] for row in detail[1:]} == {
        "Septembre 2026 (15.09 → 30.09)", "Octobre 2026 (01.10 → 15.10)",
    }
    sept = _par_zone(_section(_csv(client_user, "mois=2026-09"), "Électricité"), 2)
    partiel = _par_zone(_section(plage, "Électricité"), 2)
    assert 0 < partiel["Communs"] < sept["Communs"] + 200
    res = client_user.get(URL.format("csv") + "&du=2026-09-15&au=2026-10-15")
    assert "decompte-Ref-2026-09-15_2026-10-15.csv" in res.headers["Content-Disposition"]


def test_classeur_avec_onglet_detail(client_user):
    res = client_user.get(URL.format("xlsx") + "&du=2026-09-01&au=2026-10-31")
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.get_data()))
    assert wb.sheetnames[:2] == ["Électricité", "Détail par mois"]
    notes = " ".join(str(c.value) for row in wb["Électricité"].iter_rows() for c in row if c.value)
    assert "Selon le mois, ces parts sont calculées" in notes
    un_mois = load_workbook(io.BytesIO(client_user.get(URL.format("xlsx") + "&mois=2026-10").get_data()))
    assert "Détail par mois" not in un_mois.sheetnames


@pytest.mark.parametrize("query, message", [
    ("du=2026-09-01", "préciser soit mois"),
    ("mois=2026-09&du=2026-09-01&au=2026-09-30", "préciser soit mois"),
    ("du=2026-09-31&au=2026-10-31", "date invalide"),
    ("du=15.09.2026&au=2026-10-31", "date invalide"),
    ("du=2026-10-31&au=2026-09-01", "la date de fin précède"),
    ("du=2020-01-01&au=2026-10-31", "limitée à 36 mois"),
    ("du=2025-01-01&au=2025-03-31", "aucune donnée"),
])
def test_periode_invalide_400(client_user, query, message):
    for ext in ("csv", "xlsx"):
        res = client_user.get(URL.format(ext) + "&" + query)
        assert res.status_code == 400
        assert message in res.get_json()["error"]


def test_la_page_lit_une_plage_choisie(client_user):
    data = client_user.get(f"/api/decompte?miniserver={ref.SITE}&du=2026-09-15&au=2026-10-15").get_json()
    cle = data["periode_cle"]
    assert [p["key"] for p in data["periodes"]] == [cle]
    assert data["periodes"][0]["label"] == "15.09.2026 → 15.10.2026"
    assert data["repartition"][cle]["source"] == "mixte"
    zones = {z["label"]: z["periodes"][cle] for z in data["zones"]}
    csv = _par_zone(_section(_csv(client_user, "du=2026-09-15&au=2026-10-15"), "Électricité"), 2)
    assert round(zones["App 1"]["total"], 2) == csv["App 1"]
    assert "postes" not in data
    assert client_user.get(f"/api/decompte?miniserver={ref.SITE}&du=2026-10-31&au=2026-09-01").status_code == 400
