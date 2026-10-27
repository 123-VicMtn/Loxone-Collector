"""
Décompte de référence : fige le comportement actuel avant le découpage du
code (docs/analyse-maintenabilite-2026-10.md, étape 2). Un déplacement de
code ne doit changer aucun kWh, aucun libellé, aucune cellule.

Sur le site synthétique « Ref » (tests/site_reference.py) : la réponse
complète de /api/decompte, le CSV et le contenu de chaque onglet du
classeur, pour septembre (séries du Miniserver) et octobre (répartition
recalculée, compteur remplacé).

Mettre à jour la référence après un changement VOULU du résultat :
    UPDATE_SNAPSHOTS=1 python -m pytest tests/test_reference_decompte.py
puis relire le diff de tests/snapshots/reference_decompte.json avant de
commiter.
"""
import hashlib
import io
import json
import os
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

import site_reference as ref

SNAPSHOT = Path(__file__).parent / "snapshots" / "reference_decompte.json"
MOIS = ("2026-09", "2026-10")


def _cellule(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _classeur(contenu: bytes) -> dict:
    wb = load_workbook(io.BytesIO(contenu))
    out = {}
    for ws in wb.worksheets:
        lignes = [[_cellule(c.value) for c in row] for row in ws.iter_rows()]
        if ws.title == "Relevés":
            # Un point horaire par ligne (~1 500 lignes) : l'empreinte suffit
            # à détecter un changement, le début reste lisible dans un diff.
            empreinte = hashlib.sha256(json.dumps(lignes, ensure_ascii=False).encode()).hexdigest()
            out[ws.title] = {"lignes": len(lignes), "debut": lignes[:5], "sha256": empreinte}
        else:
            out[ws.title] = lignes
    return out


def _resultat_actuel(client) -> dict:
    out = {"decompte": client.get(f"/api/decompte?miniserver={ref.SITE}").get_json()}
    for mois in MOIS:
        csv = client.get(f"/api/decompte/export.csv?miniserver={ref.SITE}&mois={mois}")
        assert csv.status_code == 200
        out[f"csv {mois}"] = csv.get_data().decode("utf-8-sig").splitlines()
        xlsx = client.get(f"/api/decompte/export.xlsx?miniserver={ref.SITE}&mois={mois}")
        assert xlsx.status_code == 200
        out[f"xlsx {mois}"] = _classeur(xlsx.get_data())
    # Même passage par JSON que la référence enregistrée (tuples -> listes).
    return json.loads(json.dumps(out, ensure_ascii=False))


def test_decompte_identique_a_la_reference(client_admin):
    actuel = _resultat_actuel(client_admin)
    if os.environ.get("UPDATE_SNAPSHOTS"):
        SNAPSHOT.parent.mkdir(exist_ok=True)
        SNAPSHOT.write_text(json.dumps(actuel, ensure_ascii=False, indent=1, sort_keys=True) + "\n")
        pytest.skip("référence réécrite")
    attendu = json.loads(SNAPSHOT.read_text())
    for cle in attendu:
        assert actuel.get(cle) == attendu[cle], f"« {cle} » diffère de la référence"
    assert set(actuel) == set(attendu)


def test_reference_couvre_les_chemins_voulus(client_admin):
    # Garde-fou sur le jeu de données lui-même : s'il ne passait plus par
    # ces chemins, la référence ne protégerait plus rien.
    rep = client_admin.get(f"/api/decompte?miniserver={ref.SITE}").get_json()["repartition"]
    assert rep["2026-09"]["source"] == "loxone" and rep["2026-09"]["confiance_verifiee"]
    assert rep["2026-10"]["source"] == "calculee"
