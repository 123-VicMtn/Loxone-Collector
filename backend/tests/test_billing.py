"""
Tests pour billing.py (page de décompte de charges, /decompte).

Quatre choses valent d'être verrouillées par un test ici :
  - le découpage MENSUEL en heure LOCALE (Europe/Zurich), qui diffère du
    reste du dashboard (découpage UTC) et se casserait silencieusement si
    quelqu'un le repassait en UTC ;
  - la distinction entre taux d'autoproduction et taux d'autoconsommation,
    qui évoluent en sens inverse et ont déjà été confondus une fois ;
  - la résolution "quelle série alimente quelle colonne", qui repose sur des
    heuristiques de libellé validées sur l'installation MS-Arlopi (le cas
    "Communs Sol", mal classé automatiquement, est le plus fragile) ;
  - le refus de calculer une consommation quand un compteur a été remis à
    zéro dans la période : c'est là qu'un décompte deviendrait faux sans
    prévenir.

Utilise unittest (stdlib), comme test_classification.py.
"""
import io
import sqlite3
import unittest
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from loxone_collector import db
from loxone_collector.exports import decompte as decompte_export
from loxone_collector.exports import decompte_xlsx
from loxone_collector.services import billing


def ts(year, month, day, hour=0, tz="Europe/Zurich"):
    return int(datetime(year, month, day, hour, tzinfo=ZoneInfo(tz)).timestamp())


class TestPeriodes(unittest.TestCase):
    def test_libelles_mensuels(self):
        self.assertEqual(billing.period_label(2026, 1), "Janvier 2026")
        self.assertEqual(billing.period_label(2026, 8), "Août 2026")
        self.assertEqual(billing.period_label_short(2026, 5), "mai. 26")

    def test_bornes_calees_sur_minuit_local(self):
        start, end = billing.period_bounds(2026, 1)
        self.assertEqual(start, ts(2026, 1, 1))
        self.assertEqual(end, ts(2026, 2, 1))

    def test_decembre_deborde_sur_lannee_suivante(self):
        start, end = billing.period_bounds(2025, 12)
        self.assertEqual(start, ts(2025, 12, 1))
        self.assertEqual(end, ts(2026, 1, 1))

    def test_bornes_pas_en_utc(self):
        # Régression : en heure suisse (UTC+1 en hiver), le 1er janvier
        # commence une heure avant le minuit UTC.
        start, _ = billing.period_bounds(2026, 1)
        utc_midnight = int(datetime(2026, 1, 1, tzinfo=ZoneInfo("UTC")).timestamp())
        self.assertEqual(utc_midnight - start, 3600)

    def test_mois_a_cheval_sur_le_changement_dheure(self):
        # Mars contient le passage à l'heure d'été : le mois dure une heure
        # de moins que 31 jours pleins.
        start, end = billing.period_bounds(2026, 3)
        self.assertEqual(end - start, 31 * 86400 - 3600)

    def test_fevrier_non_bissextile(self):
        start, end = billing.period_bounds(2026, 2)
        self.assertEqual((end - start) // 86400, 28)

    def test_periodes_couvrantes(self):
        keys = [p["key"] for p in billing.periods_covering(ts(2025, 11, 15), ts(2026, 2, 2))]
        self.assertEqual(keys, ["2025-11", "2025-12", "2026-01", "2026-02"])

    def test_cle_de_periode_aller_retour(self):
        self.assertEqual(billing.parse_period_key("2026-03"), (2026, 3))
        with self.assertRaises(ValueError):
            billing.parse_period_key("2026-13")
        with self.assertRaises(ValueError):
            billing.parse_period_key("2026-P1")


def s(series_id, label, resource_type, apartment, state_name="total"):
    return {
        "series_id": series_id, "label": label, "resource_type": resource_type,
        "apartment": apartment, "state_name": state_name, "unit": "kWh",
    }


class TestResolutionDesSeries(unittest.TestCase):
    """Cas réels observés sur MS-Arlopi (voir CLAUDE.md)."""

    def test_zone_standard(self):
        series = [
            s("g", "App 1 Grid (total)", "energie_reseau", "APP1"),
            s("s", "App 1 Solaire (total)", "energie_solaire", "APP1"),
            s("c", "Appartement 1 (total)", "energie_consommee", "APP1"),
        ]
        z = billing.resolve_zones(series)[0]
        self.assertEqual(z["sources"]["reseau"]["series_id"], "g")
        self.assertEqual(z["sources"]["solaire"]["series_id"], "s")
        self.assertEqual(z["sources"]["controle"]["series_id"], "c")

    def test_le_chauffage_nest_pas_le_compteur_de_controle(self):
        # Périmètre séparé, vérifié sur les données.
        series = [
            s("c", "Appartement 1 (total)", "energie_consommee", "APP1"),
            s("ch", "Chauffage App 1 (total)", "energie_consommee", "APP1"),
        ]
        z = billing.resolve_zones(series)[0]
        self.assertEqual(z["sources"]["controle"]["series_id"], "c")

    def test_communs_sol_reconnu_comme_solaire_malgre_sa_classification(self):
        # "Communs Sol" est rangé en "énergie consommée" par
        # classification.py (pas de mot-clé solaire reconnu dans "Sol") :
        # le repli sur le libellé doit le rattraper, et surtout ne pas le
        # confondre avec le compteur de contrôle de la zone.
        series = [
            s("g", "Communs Grid (total)", "energie_reseau", "COMMUN"),
            s("sol", "Communs Sol (total)", "energie_consommee", "COMMUN"),
            s("c", "Commun et Boiler (total)", "energie_consommee", "COMMUN"),
        ]
        z = billing.resolve_zones(series)[0]
        self.assertEqual(z["sources"]["solaire"]["series_id"], "sol")
        self.assertEqual(z["sources"]["controle"]["series_id"], "c")

    def test_seuls_les_compteurs_cumulatifs_sont_retenus(self):
        # "actual" est une puissance instantanée en kW : en faire une
        # différence de relevés n'aurait aucun sens.
        series = [
            s("a", "App 1 Grid (actual)", "energie_reseau", "APP1", state_name="actual"),
            s("g", "App 1 Grid (total)", "energie_reseau", "APP1"),
        ]
        z = billing.resolve_zones(series)[0]
        self.assertEqual(z["sources"]["reseau"]["series_id"], "g")

    def test_colonne_sans_serie(self):
        series = [s("c", "Rez jardin (total)", "energie_consommee", "REZJARDIN")]
        z = billing.resolve_zones(series)[0]
        self.assertIsNone(z["sources"]["reseau"])
        self.assertEqual(z["label"], "Rez Jardin")

    def test_wallbox_est_une_zone_a_part(self):
        # Borne de recharge : un compteur d'énergie dédié, sans appartement
        # dans Loxone Config -- reclassé manuellement comme Salles Communes.
        series = [s("w", "Wallbox Energizähler (total)", "energie_consommee", "WALLBOX")]
        z = billing.resolve_zones(series)[0]
        self.assertEqual(z["zone"], "WALLBOX")
        self.assertEqual(z["label"], "Wallbox")
        self.assertEqual(z["sources"]["controle"]["series_id"], "w")
        self.assertIsNone(z["sources"]["reseau"])

    def test_batiment(self):
        series = [
            s("p", "Production (total)", "energie_consommee", ""),
            s("i", "Réseau (total)", "energie_reseau", ""),
            s("e", "Réseau (totalNeg)", "energie_reseau", "", state_name="totalNeg"),
            s("z", "Appartement 1 (total)", "energie_consommee", "APP1"),
        ]
        b = billing.resolve_batiment(series)
        self.assertEqual(b["production"]["series_id"], "p")
        self.assertEqual(b["reseau_import"]["series_id"], "i")
        self.assertEqual(b["reseau_export"]["series_id"], "e")

    def test_solaire_total_prime_sur_production_figee(self):
        # Sequoia : « Production Solaire » ne bouge plus, « Solaire (total) » si.
        series = [
            s("fige", "Production Solaire (total)", "energie_solaire", ""),
            s("vif", "Solaire (total)", "energie_solaire", ""),
            s("i", "Réseau (total)", "energie_reseau", ""),
            s("e", "Réseau (totalNeg)", "energie_reseau", "", state_name="totalNeg"),
        ]
        self.assertEqual(billing.resolve_batiment(series)["production"]["series_id"], "vif")

    def test_fenetre_passee_prend_le_compteur_qui_a_des_heures(self):
        # « Solaire (total) » n'existe qu'après le mois demandé.
        conn = sqlite3.connect(":memory:")
        conn.executescript(db.SCHEMA)
        start, end = billing.period_bounds(2026, 8)
        conn.executemany(
            "INSERT INTO readings_hourly (series_id, ts, avg_value, max_value) VALUES (?, ?, ?, ?)",
            [
                ("hist", start, 100.0, 100.0),
                ("hist", start + 3600, 110.0, 110.0),
                ("vif", ts(2026, 9, 24), 1.0, 1.0),
            ],
        )
        series = [
            s("hist", "Production Solaire (total)", "energie_solaire", ""),
            s("vif", "Solaire (total)", "energie_solaire", ""),
        ]
        chosen = billing.production_couvrante(conn, series, start, end)
        self.assertEqual(chosen["series_id"], "hist")
        conn.close()

    def test_solaire_et_batterie_ne_remplace_pas_la_production(self):
        # Arlopi : « Solaire & Batterie » n'est pas le compteur de production.
        series = [
            s("bat", "Solaire & Batterie (total)", "energie_batterie", ""),
            s("p", "Production (total)", "energie_consommee", ""),
        ]
        self.assertEqual(billing.resolve_batiment(series)["production"]["series_id"], "p")

    def test_batiment_convention_la_calbasse(self):
        series = [
            s("pv", "Compteur PV (total)", "energie_solaire", ""),
            s("i", "Compteur Intro (total)", "energie_consommee", ""),
            s("e", "Compteur Intro (totalNeg)", "energie_consommee", "", state_name="totalNeg"),
        ]
        b = billing.resolve_batiment(series)
        self.assertEqual(b["production"]["series_id"], "pv")
        self.assertEqual(b["reseau_import"]["series_id"], "i")
        self.assertEqual(b["reseau_export"]["series_id"], "e")

    def test_compteurs_pv_dhorizon_ne_sont_pas_la_production(self):
        series = [
            s("ppe", "Compteur PV PPE (total)", "energie_solaire", ""),
            s("gar", "Compteur PV garages (total)", "energie_solaire", ""),
        ]
        self.assertIsNone(billing.resolve_batiment(series)["production"])

    def test_pac_du_batiment_est_une_zone_a_part(self):
        # La Calbasse : la PAC n'a pas de zone dans Loxone Config, mais a ses
        # parts réseau / solaire, hors lots et hors Communs.
        series = [
            s("pac", "PAC (total)", "energie_consommee", ""),
            s("pr", "PAC Réseau (total)", "energie_reseau", ""),
            s("ps", "PAC Solaire (total)", "energie_solaire", ""),
            s("c", "Commun (total)", "energie_consommee", "COMMUN"),
        ]
        zones = {z["zone"]: z for z in billing.resolve_zones(series)}
        self.assertEqual(list(zones), ["COMMUN", "PAC"])
        pac = zones["PAC"]
        self.assertEqual(pac["label"], "PAC")
        self.assertEqual(pac["sources"]["controle"]["series_id"], "pac")
        self.assertEqual(pac["sources"]["reseau"]["series_id"], "pr")
        self.assertEqual(pac["sources"]["solaire"]["series_id"], "ps")

    def test_pac_affichee_sous_reserve(self):
        # Sous-compteur des Communs ou non : à confirmer sur place
        # (utilisateur, 2026-10-27). La réserve suit le payload et le classeur.
        conn = sqlite3.connect(":memory:")
        conn.executescript(db.SCHEMA)
        series = [
            s("pac", "PAC (total)", "energie_consommee", ""),
            s("c", "Commun (total)", "energie_consommee", "COMMUN"),
        ]
        payload = billing.compute_decompte(conn, series, [], [], ts(2026, 10, 27))
        conn.close()
        self.assertEqual(payload["reserves"], [billing.RESERVES_ZONES["PAC"]])
        notes = [texte for texte, _ in decompte_xlsx._notes_classeur(payload, "2026-10", [])]
        self.assertIn("Sous réserve", notes)
        self.assertIn(billing.RESERVES_ZONES["PAC"], notes)

    def test_pas_de_reserve_sans_pac(self):
        conn = sqlite3.connect(":memory:")
        conn.executescript(db.SCHEMA)
        series = [s("c", "Commun (total)", "energie_consommee", "COMMUN")]
        payload = billing.compute_decompte(conn, series, [], [], ts(2026, 10, 27))
        conn.close()
        self.assertEqual(payload["reserves"], [])

    def test_pac_deja_mesuree_de_sequoia_reste_hors_zone(self):
        # Comprise dans le compteur des Communs : en faire une zone la
        # compterait deux fois.
        series = [s("pac", "PAC déjà mesuré (total)", "energie_consommee", "")]
        self.assertEqual(billing.resolve_zones(series), [])


class TestTaux(unittest.TestCase):
    """Les deux taux évoluent en sens INVERSE au fil des saisons. Les avoir
    confondus une fois a produit un affichage qui semblait inversé (taux
    élevé en hiver) -- ce test fige la distinction."""

    def test_autoproduction_monte_en_ete(self):
        # Février : 19 kWh de solaire pour 1733 kWh consommés.
        hiver = billing.taux(19.0, 1733.0)
        # Juillet : 454 kWh de solaire pour 1950 kWh consommés.
        ete = billing.taux(454.0, 1950.0)
        self.assertLess(hiver, 2.0)
        self.assertGreater(ete, 20.0)
        self.assertGreater(ete, hiver)

    def test_autoconsommation_baisse_en_ete(self):
        # Même solaire autoconsommé, rapporté cette fois à la PRODUCTION.
        hiver = billing.taux(19.0, 23.0)      # février : on produit 23 kWh
        ete = billing.taux(454.0, 1578.0)     # juillet : on produit 1578 kWh
        self.assertGreater(hiver, 80.0)
        self.assertLess(ete, 35.0)
        self.assertLess(ete, hiver)

    def test_denominateur_nul_donne_none_pas_zero(self):
        # "aucune production du mois" et "production entièrement réinjectée"
        # ne sont pas la même information.
        self.assertIsNone(billing.taux(0.0, 0.0))
        self.assertIsNone(billing.taux(None, 100.0))


class TestConsommationDunePeriode(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        self.conn.executescript(db.SCHEMA)

    def tearDown(self):
        self.conn.close()

    def hourly(self, series_id, points):
        self.conn.executemany(
            "INSERT INTO readings_hourly (series_id, ts, avg_value) VALUES (?, ?, ?)",
            [(series_id, t, v) for t, v in points],
        )
        self.conn.commit()

    def test_delta_entre_deux_releves(self):
        self.hourly("m", [(ts(2026, 1, 1), 100.0), (ts(2026, 1, 15), 160.0),
                          (ts(2026, 1, 31, 23), 180.0), (ts(2026, 2, 5), 190.0)])
        start, end = billing.period_bounds(2026, 1)
        res = billing.reading_delta(self.conn, "m", start, end, ts(2026, 6, 1))
        # Le relevé du 5 février appartient au mois SUIVANT : le compter
        # ici le ferait apparaître deux fois sur deux factures.
        self.assertAlmostEqual(res["kwh"], 80.0)
        self.assertEqual(res["alertes"], [])

    def test_releve_anterieur_a_la_periode_sert_de_point_de_depart(self):
        self.hourly("m", [(ts(2025, 12, 31, 23), 100.0), (ts(2026, 1, 20), 150.0)])
        start, end = billing.period_bounds(2026, 1)
        res = billing.reading_delta(self.conn, "m", start, end, ts(2026, 6, 1))
        self.assertAlmostEqual(res["kwh"], 50.0)

    def test_reset_de_compteur_rend_la_periode_non_calculable(self):
        # Cas réel : le compteur "Réseau" de MS-Arlopi est passé de 9795 à
        # 1344 le 26.03.2026. Un delta brut donnerait -8451 kWh, une somme
        # de deltas positifs donnerait un chiffre plausible mais faux.
        self.hourly("m", [(ts(2025, 12, 31, 23), 90.0),
                          (ts(2026, 1, 5), 100.0), (ts(2026, 1, 20), 9795.0),
                          (ts(2026, 1, 22), 1344.0), (ts(2026, 1, 25), 1500.0)])
        start, end = billing.period_bounds(2026, 1)
        res = billing.reading_delta(self.conn, "m", start, end, ts(2026, 6, 1))
        self.assertIsNone(res["kwh"])
        self.assertEqual(len(res["ruptures"]), 1)

    def test_micro_baisse_darrondi_ignoree(self):
        # Écart d'arrondi entre le poller live (pleine précision) et
        # l'historique Statistics importé : ne doit PAS passer pour un reset.
        self.hourly("m", [(ts(2025, 12, 31, 23), 2000.0),
                          (ts(2026, 1, 5), 2048.0454678), (ts(2026, 1, 20), 2048.045),
                          (ts(2026, 1, 25), 2100.0)])
        start, end = billing.period_bounds(2026, 1)
        res = billing.reading_delta(self.conn, "m", start, end, ts(2026, 6, 1))
        self.assertIsNotNone(res["kwh"])
        self.assertNotIn("ruptures", res)

    def test_trou_de_collecte_signale(self):
        self.hourly("m", [(ts(2025, 11, 1), 50.0), (ts(2026, 1, 20), 150.0)])
        start, end = billing.period_bounds(2026, 1)
        res = billing.reading_delta(self.conn, "m", start, end, ts(2026, 6, 1))
        self.assertTrue(any("trou de collecte" in a for a in res["alertes"]))

    def test_le_decompte_ignore_le_collecteur(self):
        # Le collecteur s'est arrêté le 28 : son dernier index donne 70 kWh.
        # Les Statistics du Miniserver, elles, vont jusqu'à la fin du mois.
        self.conn.execute(
            "INSERT INTO readings (series_id, ts, value) VALUES (?, ?, ?)",
            ("m", ts(2026, 9, 28, 17), 170.0),
        )
        self.conn.executemany(
            "INSERT INTO statistic_samples (series_id, ts, value) VALUES (?, ?, ?)",
            [
                ("m", ts(2026, 9, 1), 100.0),
                ("m", ts(2026, 9, 30, 23), 177.1),
            ],
        )
        self.conn.commit()
        start, end = billing.period_bounds(2026, 9)
        res = billing.reading_delta(self.conn, "m", start, end, ts(2026, 10, 2), statistics=True)
        self.assertAlmostEqual(res["kwh"], 77.1)
        self.assertEqual(res["alertes"], [])

    def test_pas_de_fausse_alerte_sur_une_periode_en_cours(self):
        # Mois en cours : la borne de fin est dans le futur, le dernier
        # relevé est forcément "vieux" par rapport à elle.
        self.hourly("m", [(ts(2026, 1, 1), 100.0), (ts(2026, 1, 20), 150.0)])
        start, end = billing.period_bounds(2026, 1)
        res = billing.reading_delta(self.conn, "m", start, end, ts(2026, 1, 20, 12))
        self.assertEqual(res["alertes"], [])


class TestTarifsEtMontants(unittest.TestCase):
    TARIFS = [
        {"valid_from": "2025-01-01", "prix_reseau": 0.25, "prix_solaire": 0.12, "taux_tva": 8.1},
        {"valid_from": "2026-01-01", "prix_reseau": 0.30, "prix_solaire": 0.15, "taux_tva": 8.1},
    ]

    def test_tarif_choisi_sur_le_debut_de_periode(self):
        # Une hausse au 01.01.2026 ne doit pas s'appliquer à décembre 2025.
        t = billing.tarif_for(self.TARIFS, billing.period_bounds(2025, 12)[0])
        self.assertEqual(t["prix_reseau"], 0.25)
        t = billing.tarif_for(self.TARIFS, billing.period_bounds(2026, 1)[0])
        self.assertEqual(t["prix_reseau"], 0.30)

    def test_aucun_tarif_avant_la_premiere_date(self):
        self.assertIsNone(billing.tarif_for(self.TARIFS, billing.period_bounds(2024, 1)[0]))

    def test_montants(self):
        m = billing.montants(100.0, 50.0, self.TARIFS[1])
        self.assertAlmostEqual(m["ht"], 100 * 0.30 + 50 * 0.15)
        self.assertAlmostEqual(m["tva"], m["ht"] * 0.081)
        self.assertAlmostEqual(m["ttc"], m["ht"] + m["tva"])

    def test_montant_absent_plutot_que_zero(self):
        # Un montant nul faute de données ne doit pas se confondre avec un
        # montant nul réellement dû.
        self.assertIsNone(billing.montants(None, 50.0, self.TARIFS[1])["ttc"])
        self.assertIsNone(billing.montants(100.0, 50.0, None)["ttc"])


class TestExportLignes(unittest.TestCase):
    def test_lignes_et_total(self):
        payload = {
            "zones": [
                {"label": "App 1", "periodes": {"2026-05": {
                    "reseau": {"kwh": 10.0}, "solaire": {"kwh": 5.0}, "total": 15.0,
                    "taux_autoproduction": 33.333, "facturable": True, "en_cours": False,
                }}},
                {"label": "App 2", "periodes": {"2026-05": {
                    "reseau": {"kwh": None}, "solaire": {"kwh": 2.0}, "total": None,
                    "taux_autoproduction": None, "facturable": False, "en_cours": False,
                }}},
            ],
        }
        rows = decompte_export.export_lignes(payload, "2026-05")
        self.assertEqual(rows[0], decompte_export.EXPORT_HEADERS)
        self.assertEqual(rows[1][0], "Mai 2026")
        self.assertEqual(rows[1][1], "App 1")
        self.assertEqual(rows[1][5], "facturable")
        self.assertEqual(rows[2][5], "données incomplètes")
        self.assertEqual(rows[3][0], "Mai 2026")
        self.assertEqual(rows[3][2], 15.0)
        self.assertEqual(rows[3][3], 7.0)
        self.assertEqual(rows[3][4], 10.0)
        self.assertEqual(rows[3][1], "Total immeuble")

    def test_mois_absent(self):
        with self.assertRaises(KeyError):
            decompte_export.export_lignes({"zones": []}, "2026-01")


class TestDecompteWorkbook(unittest.TestCase):
    """Le classeur doit s'ouvrir déjà lisible : titre, colonnes assez larges,
    nombres en kWh, prix préremplis, montants en formules, relevés dans
    un second onglet, mois en cours borné au dernier relevé."""

    TARIF = {
        "prix_reseau": 0.25, "prix_solaire": 0.12, "taux_tva": 8.1,
        "valid_from": "2026-01-01",
    }

    def setUp(self):
        self.payload = {
            "timezone": "Europe/Zurich",
            "generated_at": int(datetime(2026, 6, 2, 10, 0, tzinfo=ZoneInfo("Europe/Zurich")).timestamp()),
            "periodes": [{
                "key": "2026-05",
                "label": "Mai 2026",
                "start": billing.period_bounds(2026, 5)[0],
                "end": billing.period_bounds(2026, 5)[1],
            }],
            "repartition": {"2026-05": {
                "source": "calculee",
                "confiance_verifiee": True,
                "raisons": ["heures d'achat au réseau impossibles"],
            }},
            "zones": [
                {"label": "App 1", "sources": {
                    "controle": {"label": "Appartement 1 (total)", "series_id": "lot", "unit": "kWh"},
                }, "periodes": {"2026-05": {
                    "reseau": {"kwh": 10.0, "alertes": []},
                    "solaire": {"kwh": 5.004, "alertes": []},
                    "controle": {
                        "kwh": 15.004, "releve_debut": 100.0, "releve_fin": 115.004,
                        "releve_debut_ts": int(datetime(2026, 5, 1, 0, 10, tzinfo=ZoneInfo("Europe/Zurich")).timestamp()),
                        "releve_fin_ts": int(datetime(2026, 5, 31, 23, 50, tzinfo=ZoneInfo("Europe/Zurich")).timestamp()),
                        "alertes": [],
                    },
                    "total": 15.004,
                    "facturable": True, "en_cours": False,
                    "tarif": self.TARIF,
                }}},
                {"label": "App 2", "periodes": {"2026-05": {
                    "reseau": {"kwh": None, "alertes": ["pas de relevé dans la période"]},
                    "solaire": {"kwh": 2.0, "alertes": []},
                    "total": None,
                    "facturable": False, "en_cours": False,
                    "tarif": self.TARIF,
                }}},
            ],
        }

    def classeur(self, payload=None, mois="2026-05", site="MS-Démo", conn=None):
        buf = io.BytesIO()
        wb = decompte_xlsx.build_decompte_workbook(payload or self.payload, mois, site, conn)
        wb.save(buf)
        buf.seek(0)
        return load_workbook(buf)

    def workbook(self, payload=None, mois="2026-05", site="MS-Démo", conn=None):
        return self.classeur(payload, mois, site, conn)["Électricité"]

    def _ligne(self, ws, valeur, col=1):
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, col).value == valeur:
                return r
        self.fail(f"{valeur!r} introuvable")

    def test_onglet_frais_supplementaires_vide_et_lie_aux_zones(self):
        wb = self.classeur()
        elec, ws = wb["Électricité"], wb["Frais supplémentaires"]
        self.assertIn("MS-Démo", ws["A1"].value)
        self.assertEqual(
            [ws.cell(4, c).value for c in range(1, 7)],
            ["Zone", "Location des compteurs (CHF)", "Frais exceptionnels (CHF)",
             "Autres frais (CHF)", "Total frais (CHF)", "Remarque"],
        )
        for row, zone in ((5, "App 1"), (6, "App 2")):
            self.assertEqual(ws.cell(row, 1).value, zone)
            cible = ws.cell(row, 1).hyperlink.location
            self.assertEqual(cible, f"'Électricité'!A{self._ligne(elec, zone)}")
            for c in (2, 3, 4, 6):
                self.assertIsNone(ws.cell(row, c).value)
            self.assertEqual(ws.cell(row, 5).value, f'=IF(COUNT(B{row}:D{row})=0,"",SUM(B{row}:D{row}))')
        self.assertEqual(ws["A7"].value, "Total immeuble")
        self.assertEqual(ws["E7"].value, '=IF(COUNT(E5:E6)=0,"",SUM(E5:E6))')

    def test_titre_largeurs_et_nombres(self):
        ws = self.workbook()
        self.assertEqual(ws.title, "Électricité")
        self.assertIn("MS-Démo", ws["A1"].value)
        self.assertIn("Mai 2026", ws["A2"].value)
        self.assertIn("01.05.2026", ws["A2"].value)
        self.assertIn("31.05.2026", ws["A2"].value)
        header = self._ligne(ws, "Zone")
        self.assertEqual(
            [ws.cell(header, c).value for c in range(1, 9)],
            ["Zone", "Consommation totale (kWh)", "Solaire (kWh)", "Réseau (kWh)",
             "Réseau TTC (CHF)", "Solaire TTC (CHF)", "Total TTC (CHF)", "État"],
        )
        self.assertIsNone(ws.cell(header, 9).value)
        self.assertGreaterEqual(ws.column_dimensions["A"].width, 22)
        self.assertGreaterEqual(ws.column_dimensions["B"].width, 18)
        app1 = self._ligne(ws, "App 1")
        app2 = self._ligne(ws, "App 2")
        total = self._ligne(ws, "Total immeuble")
        self.assertEqual(ws.cell(app1, 2).value, 15.0)
        self.assertEqual(ws.cell(app1, 3).value, 5.0)
        self.assertEqual(ws.cell(app1, 4).value, 10.0)
        self.assertEqual(ws.cell(app1, 2).number_format, "#,##0.00")
        self.assertIsNone(ws.cell(app2, 4).value)
        self.assertEqual(ws.cell(app2, 8).value, "données incomplètes")
        self.assertEqual(ws.cell(total, 2).value, 15.0)
        self.assertEqual(ws.cell(total, 3).value, 7.0)
        self.assertEqual(ws.cell(total, 4).value, 10.0)
        self.assertEqual(ws.freeze_panes, f"A{header + 1}")
        self.assertEqual(ws.page_setup.orientation, "landscape")
        self.assertEqual(ws.page_setup.fitToWidth, 1)
        self.assertEqual(int(ws.page_setup.paperSize), 9)
        self.assertIn(ws.page_setup.fitToHeight, (0, None))
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertTrue(any(str(r).startswith("A1:H") for r in ws.merged_cells.ranges))

    def test_le_total_egal_la_somme_des_lignes_affichees(self):
        ws = self.workbook()
        app1 = self._ligne(ws, "App 1")
        app2 = self._ligne(ws, "App 2")
        total = self._ligne(ws, "Total immeuble")
        self.assertEqual(ws.cell(total, 3).value, ws.cell(app1, 3).value + ws.cell(app2, 3).value)
        self.assertEqual(ws.cell(total, 2).value, ws.cell(app1, 2).value)
        textes = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        self.assertTrue(any(isinstance(t, str) and "47 %" in t for t in textes))

    def test_prix_pre_remplis_et_formules_de_montant(self):
        ws = self.workbook()
        self.assertEqual(ws["B4"].value, 0.25)
        self.assertEqual(ws["B5"].value, 0.12)
        self.assertEqual(ws["B6"].value, 8.1)
        self.assertTrue(str(ws["B4"].fill.fgColor.rgb).endswith("FFF9C4"))
        self.assertIn("01.01.2026", ws["C4"].value)
        app1 = self._ligne(ws, "App 1")
        app2 = self._ligne(ws, "App 2")
        self.assertEqual(
            ws.cell(app1, 5).value,
            f'=IF(OR(D{app1}="",$B$4="",$B$6=""),"",D{app1}*$B$4*(1+$B$6/100))',
        )
        self.assertEqual(
            ws.cell(app1, 6).value,
            f'=IF(OR(C{app1}="",$B$5="",$B$6=""),"",C{app1}*$B$5*(1+$B$6/100))',
        )
        self.assertEqual(
            ws.cell(app1, 7).value,
            f'=IF(OR(E{app1}="",F{app1}=""),"",E{app1}+F{app1})',
        )
        self.assertIn('D{0}=""'.format(app2), ws.cell(app2, 5).value)
        self.assertIn("Les montants sont TTC", ws["C4"].value)
        total = self._ligne(ws, "Total immeuble")
        self.assertEqual(ws.cell(total, 7).value, f'=IF(COUNT(G{app1}:G{app2})=0,"",SUM(G{app1}:G{app2}))')

    def test_sans_tarif_les_prix_restent_vides(self):
        for z in self.payload["zones"]:
            z["periodes"]["2026-05"]["tarif"] = None
        ws = self.workbook()
        self.assertIsNone(ws["B4"].value)
        self.assertIsNone(ws["B5"].value)
        self.assertIsNone(ws["B6"].value)
        self.assertIn("Aucun tarif", ws["C4"].value)

    def test_onglet_releves(self):
        start, end = billing.period_bounds(2026, 5)
        conn = sqlite3.connect(":memory:")
        conn.executescript(db.SCHEMA)
        # Trois points dans le mois, un avant : seul le mois figure.
        for ts, value in (
            (start - 3600, 99.0),
            (start, 100.0),
            (start + 3600, 100.5),
            (end - 3600, 115.004),
        ):
            conn.execute(
                "INSERT INTO statistic_samples (series_id, ts, value) VALUES ('lot', ?, ?)",
                (ts, value),
            )
        conn.commit()
        wb = self.classeur(conn=conn)
        self.assertEqual(wb.sheetnames, ["Électricité", "Frais supplémentaires", "Relevés"])
        ws = wb["Relevés"]
        self.assertIn("Relevés de compteurs", ws["A1"].value)
        self.assertIn("Mai 2026", ws["A2"].value)
        self.assertEqual([ws.cell(4, c).value for c in range(1, 5)], ["Zone", "Compteur", "Date", "Index"])
        self.assertIsNone(ws.cell(4, 5).value)
        self.assertEqual(ws["A5"].value, "App 1")
        self.assertEqual(ws["B5"].value, "Appartement 1 (total)")
        self.assertIn("01.05.2026", ws["C5"].value)
        self.assertEqual(ws["D5"].value, 100.0)
        self.assertEqual(ws["D6"].value, 100.5)
        self.assertEqual(ws["D7"].value, 115.004)
        self.assertIsNone(ws["A8"].value)
        self.assertNotIn("99", [str(ws.cell(r, 4).value) for r in range(5, 8)])
        self.assertGreaterEqual(ws.column_dimensions["B"].width, 30)
        self.assertFalse(ws.sheet_view.showGridLines)
        conn.close()

    def test_notes_de_lecture_et_releves_manquants(self):
        ws = self.workbook()
        textes = []
        for r in range(1, ws.max_row + 1):
            for c in range(1, 4):
                v = ws.cell(r, c).value
                if isinstance(v, str):
                    textes.append(v)
        bloc = "\n".join(textes)
        self.assertIn("Ces parts sont calculées", bloc)
        self.assertNotIn("heures d'achat au réseau impossibles", bloc)
        self.assertIn("App 2 — Réseau : pas de relevé dans la période", bloc)
        self.assertIn("Une case vide n'est pas un zéro", bloc)

    def test_alerte_de_zone_quand_les_colonnes_n_en_ont_pas(self):
        self.payload["zones"].append({"label": "App 100", "periodes": {"2026-05": {
            "reseau": {"kwh": None, "alertes": []},
            "solaire": {"kwh": None, "alertes": []},
            "total": None,
            "facturable": False, "en_cours": False,
            "alertes": ["compteur du lot introuvable : répartition impossible"],
        }}})
        ws = self.workbook()
        textes = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value]
        self.assertIn(
            "App 100 : compteur du lot introuvable : répartition impossible",
            "\n".join(textes),
        )

    def test_mois_en_cours_s_arrete_au_dernier_releve(self):
        start, end = billing.period_bounds(2026, 10)
        dernier = int(datetime(2026, 10, 1, 15, 30, tzinfo=ZoneInfo("Europe/Zurich")).timestamp())
        payload = {
            "timezone": "Europe/Zurich",
            "generated_at": dernier + 60,
            "periodes": [{"key": "2026-10", "label": "Octobre 2026", "start": start, "end": end}],
            "batiment": {"periodes": {"2026-10": {"en_cours": True, "production": {}}}},
            "zones": [{"label": "App 1", "periodes": {"2026-10": {
                "reseau": {"kwh": 1.0, "alertes": [], "releve_fin_ts": dernier},
                "solaire": {"kwh": 1.0, "alertes": []},
                "total": 2.0,
                "facturable": False, "en_cours": True,
            }}}],
        }
        ws = self.workbook(payload, "2026-10", "MS-PPE-Sequoia")
        self.assertIn("01.10.2026", ws["A2"].value)
        self.assertIn("dernier relevé", ws["A2"].value)
        self.assertNotIn("31.10.2026", ws["A2"].value)

    def test_mois_absent(self):
        with self.assertRaises(KeyError):
            decompte_xlsx.build_decompte_workbook({"zones": []}, "2026-01", "site")


class TestRepartitionDansLeDecompte(unittest.TestCase):
    """Quand evaluer_site condamne les séries du Miniserver, le décompte
    facture le compteur du lot, scindé. Sinon il garde Grid + Solaire."""

    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        self.conn.executescript(db.SCHEMA)

    def tearDown(self):
        self.conn.close()

    def index(self, series_id, points):
        self.conn.executemany(
            "INSERT INTO statistic_samples (series_id, ts, value) VALUES (?, ?, ?)",
            [(series_id, t, v) for t, v in points],
        )
        self.conn.commit()

    def test_series_incoherentes_sont_remplacees_par_le_compteur_du_lot(self):
        # Heure 1 : le bâtiment n'importe rien, le Grid du lot augmente quand
        # même. Heure 2 : mix normal. Le compteur du lot suit le bâtiment.
        t0 = ts(2026, 1, 1)
        self.index("prod", [(t0, 0.0), (t0 + 3600, 10.0), (t0 + 7200, 16.0)])
        self.index("imp", [(t0, 0.0), (t0 + 3600, 0.0), (t0 + 7200, 4.0)])
        self.index("exp", [(t0, 0.0), (t0 + 3600, 2.0), (t0 + 7200, 3.0)])
        self.index("lot", [(t0, 0.0), (t0 + 3600, 8.0), (t0 + 7200, 17.0)])
        self.index("grid", [(t0, 0.0), (t0 + 3600, 5.0), (t0 + 7200, 6.0)])
        self.index("sol", [(t0, 0.0), (t0 + 3600, 0.0), (t0 + 7200, 4.0)])
        series = [
            s("prod", "Solaire (total)", "energie_solaire", ""),
            s("fige", "Production Solaire (total)", "energie_solaire", ""),
            s("imp", "Réseau (total)", "energie_reseau", ""),
            s("exp", "Réseau (totalNeg)", "energie_reseau", "", state_name="totalNeg"),
            s("lot", "Appartement 1 (total)", "energie_consommee", "APP1"),
            s("grid", "App 1 Grid (total)", "energie_reseau", "APP1"),
            s("sol", "App 1 Solaire (total)", "energie_solaire", "APP1"),
        ]
        payload = billing.compute_decompte(
            self.conn, series, billing.periods_covering(t0, t0 + 7200), [], ts(2026, 6, 1),
        )
        self.assertEqual(payload["repartition"]["2026-01"]["source"], "calculee")
        self.assertEqual(payload["batiment"]["sources"]["production"]["series_id"], "prod")
        e = payload["zones"][0]["periodes"]["2026-01"]
        self.assertEqual(e["reseau"]["kwh"], 4.0)
        self.assertEqual(e["solaire"]["kwh"], 13.0)
        self.assertEqual(e["total"], 17.0)
        self.assertEqual(e["reseau"]["kwh"] + e["solaire"]["kwh"], e["total"])
        self.assertTrue(e["facturable"])
        self.assertEqual(e["reseau"]["alertes"], [])

    def test_bilan_non_exploitable_garde_les_series_du_miniserver(self):
        # Le compteur réseau du bâtiment ne mesure pas la même chose que
        # les zones : on ne recalcule pas, on garde Grid + Solaire.
        t0 = ts(2026, 1, 1)
        self.index("prod", [(t0, 0.0), (t0 + 3600, 1.0)])
        self.index("imp", [(t0, 0.0), (t0 + 3600, 1.0)])
        self.index("exp", [(t0, 0.0), (t0 + 3600, 0.0)])
        self.index("lot", [(t0, 0.0), (t0 + 3600, 100.0)])
        self.index("grid", [(t0, 0.0), (t0 + 3600, 40.0)])
        self.index("sol", [(t0, 0.0), (t0 + 3600, 60.0)])
        series = [
            s("prod", "Production (total)", "energie_consommee", ""),
            s("imp", "Réseau (total)", "energie_reseau", ""),
            s("exp", "Réseau (totalNeg)", "energie_reseau", "", state_name="totalNeg"),
            s("lot", "Appartement 1 (total)", "energie_consommee", "APP1"),
            s("grid", "App 1 Grid (total)", "energie_reseau", "APP1"),
            s("sol", "App 1 Solaire (total)", "energie_solaire", "APP1"),
        ]
        payload = billing.compute_decompte(
            self.conn, series, billing.periods_covering(t0, t0 + 3600), [], ts(2026, 6, 1),
        )
        self.assertEqual(payload["repartition"]["2026-01"]["source"], "loxone")
        e = payload["zones"][0]["periodes"]["2026-01"]
        self.assertAlmostEqual(e["reseau"]["kwh"], 40.0)
        self.assertAlmostEqual(e["solaire"]["kwh"], 60.0)
        self.assertAlmostEqual(e["total"], 100.0)

    def test_verdict_propre_a_chaque_mois(self):
        # Janvier : le Grid du lot augmente sans import du bâtiment (bloc
        # faux, comme Sequoia avant le 10.09.2026). Février : séries justes.
        # Le mois faux ne doit pas condamner le mois juste.
        t0, f0 = ts(2026, 1, 1), ts(2026, 2, 1)
        self.index("prod", [(t0, 0.0), (t0 + 3600, 10.0), (t0 + 7200, 16.0), (f0, 16.0), (f0 + 3600, 32.0)])
        self.index("imp", [(t0, 0.0), (t0 + 3600, 0.0), (t0 + 7200, 4.0), (f0, 4.0), (f0 + 3600, 8.0)])
        self.index("exp", [(t0, 0.0), (t0 + 3600, 2.0), (t0 + 7200, 3.0), (f0, 3.0), (f0 + 3600, 6.0)])
        self.index("lot", [(t0, 0.0), (t0 + 3600, 8.0), (t0 + 7200, 17.0), (f0, 17.0), (f0 + 3600, 34.0)])
        self.index("grid", [(t0, 0.0), (t0 + 3600, 5.0), (t0 + 7200, 6.0), (f0, 6.0), (f0 + 3600, 10.0)])
        self.index("sol", [(t0, 0.0), (t0 + 3600, 0.0), (t0 + 7200, 4.0), (f0, 4.0), (f0 + 3600, 17.0)])
        series = [
            s("prod", "Solaire (total)", "energie_solaire", ""),
            s("imp", "Réseau (total)", "energie_reseau", ""),
            s("exp", "Réseau (totalNeg)", "energie_reseau", "", state_name="totalNeg"),
            s("lot", "Appartement 1 (total)", "energie_consommee", "APP1"),
            s("grid", "App 1 Grid (total)", "energie_reseau", "APP1"),
            s("sol", "App 1 Solaire (total)", "energie_solaire", "APP1"),
        ]
        payload = billing.compute_decompte(
            self.conn, series, billing.periods_covering(t0, f0 + 3600), [], ts(2026, 6, 1),
        )
        janvier, fevrier = payload["repartition"]["2026-01"], payload["repartition"]["2026-02"]
        self.assertEqual(janvier["source"], "calculee")
        self.assertTrue(janvier["raisons"])
        self.assertEqual(fevrier["source"], "loxone")
        self.assertTrue(fevrier["confiance_verifiee"])
        e = payload["zones"][0]["periodes"]["2026-02"]
        self.assertAlmostEqual(e["reseau"]["kwh"], 4.0)
        self.assertAlmostEqual(e["solaire"]["kwh"], 13.0)
        self.assertEqual(billing.repartition_du_mois(payload, "2026-02"), fevrier)
        self.assertEqual(billing.repartition_du_mois(payload, "2030-01")["source"], "loxone")

    def test_releve_de_compteurs_qui_se_relaient(self):
        # « Production Solaire » couvre le début du mois, « Solaire (total) »
        # la fin. Aucun des deux seul ne ferme le bilan ; ensemble, le
        # décompte reste sur le compteur du lot.
        t0 = ts(2026, 9, 1)
        self.index("hist", [(t0, 1000.0), (t0 + 3600, 1010.0)])
        self.index("vif", [(t0 + 7200, 0.0), (t0 + 10800, 6.0)])
        self.index("imp", [
            (t0, 0.0), (t0 + 3600, 0.0), (t0 + 7200, 0.0), (t0 + 10800, 4.0),
        ])
        self.index("exp", [
            (t0, 0.0), (t0 + 3600, 2.0), (t0 + 7200, 2.0), (t0 + 10800, 3.0),
        ])
        self.index("lot", [(t0, 0.0), (t0 + 3600, 8.0), (t0 + 10800, 17.0)])
        self.index("grid", [(t0, 0.0), (t0 + 3600, 5.0), (t0 + 10800, 6.0)])
        self.index("sol", [(t0, 0.0), (t0 + 3600, 0.0), (t0 + 10800, 4.0)])
        series = [
            s("hist", "Production Solaire (total)", "energie_solaire", ""),
            s("vif", "Solaire (total)", "energie_solaire", ""),
            s("imp", "Réseau Général (total)", "energie_reseau", ""),
            s("exp", "Réseau Général (totalNeg)", "energie_reseau", "", state_name="totalNeg"),
            s("lot", "Appartement 1 (total)", "energie_consommee", "APP1"),
            s("grid", "App 1 Grid (total)", "energie_reseau", "APP1"),
            s("sol", "App 1 Solaire (total)", "energie_solaire", "APP1"),
        ]
        payload = billing.compute_decompte(
            self.conn, series, billing.periods_covering(t0, t0 + 10800), [], ts(2026, 10, 2),
        )
        self.assertEqual(payload["repartition"]["2026-09"]["source"], "calculee")
        e = payload["zones"][0]["periodes"]["2026-09"]
        self.assertEqual(e["total"], 17.0)
        self.assertEqual(e["reseau"]["kwh"] + e["solaire"]["kwh"], e["total"])
        prod = payload["batiment"]["periodes"]["2026-09"]["production"]["kwh"]
        self.assertAlmostEqual(prod, 16.0)


class TestBatimentPartiel(unittest.TestCase):
    def test_la_consommation_reste_quand_une_zone_manque(self):
        zones = [
            {"periodes": {"2026-08": {
                "reseau": {"kwh": 10.0}, "solaire": {"kwh": 5.0}, "total": 15.0,
            }}},
            {"periodes": {"2026-08": {
                "reseau": {"kwh": None}, "solaire": {"kwh": 2.0}, "total": None,
            }}},
        ]
        p = {"key": "2026-08", "start": 0, "end": 1}
        b = billing._batiment_period(None, {"production": None}, zones, p, 0)
        self.assertEqual(b["consommation_totale"], 15.0)
        self.assertEqual(b["autoconsommation"], 5.0)
        self.assertTrue(b["zones_incompletes"])
        self.assertIsNone(b["injection"])


class TestPostes(unittest.TestCase):
    """Chauffage (énergie seulement), eau chaude (pas l'eau froide) et
    buanderie, sur les libellés réels de Sequoia et d'Arlopi."""

    SERIES = [
        s("ch", "APP 01 compteur chaleur [5] (total)", "energie_consommee", "APP01"),
        s("chv", "APP 01 compteur chaleur volume [5] (total)", "energie_consommee", "APP01"),
        s("ec", "APP 01 compteur eau chaude m^3 [6] (total)", "eau_chaude", "APP01"),
        s("ef", "APP 01 compteur eau m^3 [7] (total)", "eau_froide", "APP01"),
        s("bk", "CEnergieApp 01 (total)", "energie_consommee", "APP01"),
        s("bc", "CNrMachineApp 01 (total)", "energie_consommee", "APP01"),
        s("lot", "Appartement 1 (total)", "energie_consommee", "APP01"),
        s("c34v", "APP 34 compteur Ccaleur volume [53] (total)", "energie_consommee", "APP34"),
        s("c34", "APP 34 compteur chaleur [53] (total)", "energie_consommee", "APP34"),
        s("b100", "CEnergieApp100 (total)", "energie_consommee", "APP100"),
        s("boiler", "Boiler ecs mesuré aussi ? (total)", "eau_chaude", ""),
        s("com1", "Chauffage Commerce (total)", "energie_consommee", "COMMERCE"),
        s("com2", "Chauffage Commerce supp. (total)", "energie_consommee", "COMMERCE"),
        s("live", "APP 01 compteur chaleur [5]", "energie_consommee", "APP01", state_name="actual"),
    ]

    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        self.conn.executescript(db.SCHEMA)

    def tearDown(self):
        self.conn.close()

    def index(self, series_id, points):
        self.conn.executemany(
            "INSERT INTO statistic_samples (series_id, ts, value) VALUES (?, ?, ?)",
            [(series_id, t, v) for t, v in points],
        )
        self.conn.commit()

    def postes(self):
        return {z["zone"]: z["compteurs"] for z in billing.resolve_postes(self.SERIES)}

    def ids(self, srcs):
        return [c["series_id"] for c in srcs]

    def test_energie_de_chauffage_sans_le_volume(self):
        p = self.postes()
        self.assertEqual(self.ids(p["APP01"]["chauffage"]), ["ch"])
        self.assertEqual(self.ids(p["APP34"]["chauffage"]), ["c34"])

    def test_eau_chaude_sans_eau_froide(self):
        self.assertEqual(self.ids(self.postes()["APP01"]["eau_chaude"]), ["ec"])

    def test_buanderie_et_cycles(self):
        p = self.postes()["APP01"]
        self.assertEqual(self.ids(p["buanderie"]), ["bk"])
        self.assertEqual(self.ids(p["buanderie_cycles"]), ["bc"])

    def test_zone_fantome_et_compteurs_d_immeuble_exclus(self):
        p = self.postes()
        self.assertNotIn("APP100", p)
        self.assertNotIn("", p)

    def test_tous_les_compteurs_d_une_zone_sont_gardes(self):
        self.assertEqual(self.ids(self.postes()["COMMERCE"]["chauffage"]), ["com1", "com2"])

    def test_releves_du_mois(self):
        start, end = billing.period_bounds(2026, 8)
        self.index("ch", [(start - 600, 1000.0), (end - 600, 1250.5)])
        self.index("ec", [(start - 600, 10.0), (end - 600, 10.875)])
        self.index("bk", [(start - 600, 37.0), (end - 600, 41.5)])
        self.index("bc", [(start - 600, 2.0), (end - 600, 5.0)])
        r = billing.compute_postes(
            self.conn, billing.resolve_postes(self.SERIES),
            {"key": "2026-08", "start": start, "end": end}, end + 86400,
        )
        app01 = next(x for x in r["chauffage"] if x["zone"] == "APP01")
        self.assertAlmostEqual(app01["kwh"], 250.5)
        self.assertEqual(app01["releve_debut"], 1000.0)
        self.assertAlmostEqual(r["eau_chaude"][0]["kwh"], 0.875)
        b = r["buanderie"][0]
        self.assertAlmostEqual(b["energie"]["kwh"], 4.5)
        self.assertAlmostEqual(b["cycles"]["kwh"], 3.0)
        self.assertIsNone(r["buanderie_depuis"])

    def test_buanderie_comptee_depuis_sa_mise_en_service(self):
        start, end = billing.period_bounds(2026, 7)
        self.index("bk", [(start, 0.0), (billing.BUANDERIE_VALID_FROM - 3600, 37.0), (end - 600, 38.0)])
        self.index("bc", [(start, 0.0), (end - 600, 1.0)])
        r = billing.compute_postes(
            self.conn, billing.resolve_postes(self.SERIES),
            {"key": "2026-07", "start": start, "end": end}, end + 86400,
        )
        self.assertAlmostEqual(r["buanderie"][0]["energie"]["kwh"], 1.0)
        self.assertEqual(r["buanderie_depuis"], billing.BUANDERIE_VALID_FROM)

    def test_pas_de_buanderie_avant_sa_mise_en_service(self):
        start, end = billing.period_bounds(2026, 5)
        r = billing.compute_postes(
            self.conn, billing.resolve_postes(self.SERIES),
            {"key": "2026-05", "start": start, "end": end}, end + 86400,
        )
        self.assertEqual(r["buanderie"], [])
        self.assertTrue(r["buanderie_avant_mise_en_service"])


class TestExportPostes(unittest.TestCase):
    """Classeur et CSV : un tableau électricité avec la colonne buanderie,
    puis chauffage, eau chaude et buanderie."""

    def setUp(self):
        start, end = billing.period_bounds(2026, 8)
        self.payload = {
            "timezone": "Europe/Zurich",
            "generated_at": end + 3600,
            "periodes": [{"key": "2026-08", "label": "Août 2026", "start": start, "end": end}],
            "zones": [
                {"zone": "APP01", "label": "App 1", "periodes": {"2026-08": {
                    "reseau": {"kwh": 10.0, "alertes": []},
                    "solaire": {"kwh": 5.0, "alertes": []},
                    "total": 15.0, "facturable": True, "en_cours": False,
                }}},
                {"zone": "COMMUN", "label": "Communs", "periodes": {"2026-08": {
                    "reseau": {"kwh": 20.0, "alertes": []},
                    "solaire": {"kwh": 1.0, "alertes": []},
                    "total": 21.0, "facturable": True, "en_cours": False,
                }}},
            ],
            "postes": {"2026-08": {
                "en_cours": False, "buanderie_depuis": None,
                "buanderie_avant_mise_en_service": False,
                "chauffage": [
                    {"zone": "APP01", "label": "App 1", "compteur": "APP 01 compteur chaleur [5] (total)",
                     "en_cours": False, "kwh": 250.5, "releve_debut": 1000.0, "releve_fin": 1250.5,
                     "alertes": []},
                    {"zone": "COMMUN", "label": "Communs", "compteur": "Commun - Caisson (total)",
                     "en_cours": False, "kwh": 100.0, "releve_debut": 1.0, "releve_fin": 101.0,
                     "alertes": []},
                    {"zone": "COMMUN", "label": "Communs", "compteur": "Commun - Cuisine (total)",
                     "en_cours": False, "kwh": None, "releve_debut": None, "releve_fin": None,
                     "alertes": ["pas de relevé dans la période"]},
                ],
                "eau_chaude": [
                    {"zone": "APP01", "label": "App 1", "compteur": "APP 01 compteur eau chaude m^3 [6] (total)",
                     "en_cours": False, "kwh": 0.8754, "releve_debut": 10.0, "releve_fin": 10.8754,
                     "alertes": []},
                ],
                "buanderie": [
                    {"zone": "APP01", "label": "App 1", "en_cours": False,
                     "energie": {"kwh": 4.5, "alertes": []}, "cycles": {"kwh": 3.0, "alertes": []}},
                ],
            }},
        }

    def classeur(self):
        buf = io.BytesIO()
        decompte_xlsx.build_decompte_workbook(self.payload, "2026-08", "MS-PPE-Sequoia").save(buf)
        buf.seek(0)
        return load_workbook(buf)

    def ligne(self, ws, valeur):
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value == valeur:
                return r
        self.fail(f"{valeur!r} introuvable")

    def test_onglets(self):
        self.assertEqual(
            self.classeur().sheetnames,
            ["Électricité", "Chauffage", "Eau chaude", "Buanderie",
             "Frais supplémentaires", "Relevés"],
        )

    def test_colonne_buanderie_dans_l_electricite(self):
        ws = self.classeur()["Électricité"]
        header = self.ligne(ws, "Zone")
        self.assertEqual(
            [ws.cell(header, c).value for c in range(1, 6)],
            ["Zone", "Consommation totale (kWh)", "Solaire (kWh)", "Réseau (kWh)", "Buanderie (kWh)"],
        )
        app1, communs = self.ligne(ws, "App 1"), self.ligne(ws, "Communs")
        self.assertEqual(ws.cell(app1, 5).value, 4.5)
        self.assertIsNone(ws.cell(communs, 5).value)
        self.assertEqual(ws.cell(self.ligne(ws, "Total immeuble"), 5).value, 4.5)
        # Les montants restent sur réseau et solaire, la buanderie n'y entre pas.
        self.assertEqual(
            ws.cell(app1, 6).value,
            f'=IF(OR(D{app1}="",$B$4="",$B$6=""),"",D{app1}*$B$4*(1+$B$6/100))',
        )
        self.assertEqual(ws.cell(header, 9).value, "État")

    def test_onglet_chauffage(self):
        ws = self.classeur()["Chauffage"]
        self.assertIn("MS-PPE-Sequoia", ws["A1"].value)
        self.assertEqual(
            [ws.cell(7, c).value for c in range(1, 8)],
            ["Zone", "Compteur", "Index début (kWh th)", "Index fin (kWh th)",
             "Énergie (kWh th)", "Montant TTC (CHF)", "État"],
        )
        self.assertEqual(ws["B8"].value, "APP 01 compteur chaleur [5]")
        self.assertEqual(ws["E8"].value, 250.5)
        self.assertEqual(ws["G10"].value, "données incomplètes")
        total = self.ligne(ws, "Total immeuble")
        self.assertEqual(ws.cell(total, 5).value, 350.5)
        textes = "\n".join(str(ws.cell(r, 1).value) for r in range(total, ws.max_row + 1))
        self.assertIn("volume circulé n'est pas repris", textes)
        self.assertIn("Communs — Commun - Cuisine : pas de relevé dans la période", textes)

    def test_prix_et_montants_ttc_du_chauffage(self):
        ws = self.classeur()["Chauffage"]
        self.assertEqual(ws["A4"].value, "Prix chauffage HT (CHF/kWh th)")
        self.assertEqual(ws["A5"].value, "TVA (%)")
        self.assertIsNone(ws["B4"].value)
        self.assertTrue(str(ws["B4"].fill.fgColor.rgb).endswith("FFF9C4"))
        self.assertIn("Montant TTC", ws["C4"].value)
        self.assertEqual(ws["F8"].value, '=IF(OR(E8="",$B$4="",$B$5=""),"",E8*$B$4*(1+$B$5/100))')
        total = self.ligne(ws, "Total immeuble")
        self.assertEqual(ws.cell(total, 6).value, '=IF(COUNT(F8:F10)=0,"",SUM(F8:F10))')

    def test_tva_reprise_du_tarif_electrique(self):
        self.payload["zones"][0]["periodes"]["2026-08"]["tarif"] = {
            "prix_reseau": 0.25, "prix_solaire": 0.12, "taux_tva": 8.1, "valid_from": "2026-01-01",
        }
        wb = self.classeur()
        self.assertEqual(wb["Chauffage"]["B5"].value, 8.1)
        self.assertEqual(wb["Eau chaude"]["B5"].value, 8.1)
        self.assertIsNone(wb["Eau chaude"]["B4"].value)

    def test_onglet_eau_chaude_au_litre_pres(self):
        ws = self.classeur()["Eau chaude"]
        self.assertEqual(ws["A4"].value, "Prix eau chaude HT (CHF/m³)")
        self.assertEqual(ws["E7"].value, "Volume (m³)")
        self.assertEqual(ws["E8"].value, 0.875)
        self.assertEqual(ws["E8"].number_format, "#,##0.000")
        self.assertEqual(ws["F8"].value, '=IF(OR(E8="",$B$4="",$B$5=""),"",E8*$B$4*(1+$B$5/100))')

    def test_onglet_buanderie(self):
        ws = self.classeur()["Buanderie"]
        self.assertEqual([ws.cell(4, c).value for c in range(1, 5)],
                         ["Zone", "Cycles", "Énergie (kWh)", "État"])
        self.assertEqual([ws.cell(5, c).value for c in range(1, 5)], ["App 1", 3.0, 4.5, "facturable"])

    def test_csv_en_sections(self):
        rows = decompte_export.export_csv_lignes(self.payload, "2026-08")
        titres = [r[0] for r in rows if len(r) == 1]
        self.assertEqual(titres, ["Électricité", "Chauffage", "Eau chaude", "Buanderie"])
        self.assertEqual(rows[1][-2], "Buanderie (kWh)")
        self.assertEqual(rows[2], ["Août 2026", "App 1", "15,00", "5,00", "10,00", "4,50", "facturable"])
        self.assertIn(["Août 2026", "Total immeuble", "", "", "", "350,50", ""], rows)
        self.assertIn(["Août 2026", "App 1", "APP 01 compteur eau chaude m^3 [6]",
                       "10,000", "10,875", "0,875", "facturable"], rows)
        self.assertIn(["Août 2026", "App 1", "3", "4,50", "facturable"], rows)

    def test_index_identiques_signales(self):
        postes = self.payload["postes"]["2026-08"]
        postes["eau_chaude"].append(dict(postes["eau_chaude"][0], zone="APP41", label="App 41",
                                         compteur="APP 41 compteur eau chaude m^3 [60!] (total)"))
        ws = self.classeur()["Eau chaude"]
        textes = "\n".join(str(ws.cell(r, 1).value) for r in range(1, ws.max_row + 1))
        self.assertIn("App 1 et App 41 : mêmes index", textes)

    def test_total_csv_des_valeurs_affichees(self):
        z = self.payload["zones"]
        z[0]["periodes"]["2026-08"]["total"] = 1.004
        z[1]["periodes"]["2026-08"]["total"] = 4.504
        rows = decompte_export.export_csv_lignes(self.payload, "2026-08")
        total = next(r for r in rows if r[1:2] == ["Total immeuble"])
        self.assertEqual(total[2], "5,50")

    def test_sans_postes_le_csv_reste_l_electricite(self):
        del self.payload["postes"]
        rows = decompte_export.export_csv_lignes(self.payload, "2026-08")
        self.assertEqual([r[0] for r in rows if len(r) == 1], ["Électricité"])
        self.assertNotIn("Buanderie (kWh)", rows[1])


if __name__ == "__main__":
    unittest.main()
