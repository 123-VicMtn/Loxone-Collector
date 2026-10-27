"""
Classeur Excel du décompte d'un mois : onglets Électricité, Chauffage, Eau
chaude, Buanderie, Frais supplémentaires et Relevés. Mise en forme seulement :
les lignes viennent de exports/decompte.py, les kWh de services/billing.py.
"""

from __future__ import annotations

from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.page import PageMargins

from loxone_collector.services.billing import (
    TIMEZONE,
    parse_period_key,
    period_bounds,
    period_label,
    repartition_du_mois,
)
from loxone_collector.exports.decompte import (
    _BUANDERIE_HEADER,
    _INDEX_DECIMALES,
    _POSTE_TABLES,
    _buanderie_par_zone,
    _fr,
    _postes_du_mois,
    etat_zone,
    lignes_buanderie,
    lignes_compteurs,
)


# Mise en page du classeur téléchargé. Même famille visuelle que
# scripts/export_appartement.py : le fichier doit s'ouvrir déjà lisible,
# sans élargir les colonnes ni deviner le site et le mois.
# Colonnes de l'onglet Électricité : (titre, largeur). La colonne
# buanderie n'existe que si le mois en a une, d'où des positions
# calculées et non figées.
_ELEC_COLONNES = {
    "zone": ("Zone", 28),
    "total": ("Consommation totale (kWh)", 22),
    "solaire": ("Solaire (kWh)", 18),
    "reseau": ("Réseau (kWh)", 18),
    "buanderie": (_BUANDERIE_HEADER, 18),
    "chf_reseau": ("Réseau TTC (CHF)", 18),
    "chf_solaire": ("Solaire TTC (CHF)", 18),
    "ttc": ("Total TTC (CHF)", 18),
    "etat": ("État", 24),
}
_ELEC_KWH = ("total", "solaire", "reseau", "buanderie")
_ELEC_CHF = ("chf_reseau", "chf_solaire", "ttc")
_PRIX_RESEAU_ROW = 4
_PRIX_SOLAIRE_ROW = 5
_TVA_ROW = 6
_XLSX_HEADER_ROW = 8
_KWH_FMT = "#,##0.00"
_M3_FMT = "#,##0.000"
_CYCLES_FMT = "#,##0"
_CHF_FMT = "#,##0.00"
_PRIX_FMT = "0.0000"
_TVA_FMT = "0.00"
_INDEX_FMT = "#,##0.000"
_NAVY = "1F4E78"
_THIN = Border(
    left=Side(style="thin", color="D4D4D4"),
    right=Side(style="thin", color="D4D4D4"),
    top=Side(style="thin", color="D4D4D4"),
    bottom=Side(style="thin", color="D4D4D4"),
)
_TOTAL_BORDER = Border(
    left=Side(style="thin", color="D4D4D4"),
    right=Side(style="thin", color="D4D4D4"),
    top=Side(style="medium", color=_NAVY),
    bottom=Side(style="thin", color="D4D4D4"),
)
_FILL_HEADER = PatternFill("solid", fgColor=_NAVY)
_FILL_TOTAL = PatternFill("solid", fgColor="FFF2CC")
_FILL_AMBER = PatternFill("solid", fgColor="FFF8E1")
_FILL_STRIPE = PatternFill("solid", fgColor="F4F7FB")
_FILL_EDIT = PatternFill("solid", fgColor="FFF9C4")
_FONT_TITLE = Font(bold=True, size=16, color=_NAVY)
_FONT_HEADER = Font(bold=True, color="FFFFFF", size=11)
_FONT_SUB = Font(italic=True, size=11, color="525252")
_FONT_NOTE = Font(size=10, color="525252")
_FONT_NOTE_HEAD = Font(bold=True, size=11, color=_NAVY)
_FONT_TOTAL = Font(bold=True, size=11)
_FONT_ETAT = {
    "facturable": Font(color="166534"),
    "mois en cours": Font(color="92400E"),
    "données incomplètes": Font(color="991B1B"),
}
_ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
_ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
_ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ALIGN_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _kwh_affiche(value) -> float | None:
    """Deux décimales, comme l'affichage d'un relevé. Une case vide reste
    vide : un zéro affiché serait un relevé à zéro, pas un relevé absent."""
    if value is None:
        return None
    return float(f"{float(value):.2f}")


def _index_affiche(value) -> float | None:
    """Trois décimales : assez pour vérifier index de fin − index de début,
    sans traînée binaire (594.2790000000001)."""
    if value is None:
        return None
    return float(f"{float(value):.3f}")


def _somme_affichee(values: list) -> float | None:
    nums = [v for v in values if v is not None]
    if not nums:
        return None
    return float(f"{sum(nums):.2f}")


def _jour(ts: int, tz_name: str) -> str:
    return datetime.fromtimestamp(ts, ZoneInfo(tz_name)).strftime("%d.%m.%Y")


def _horodatage(ts: int, tz_name: str) -> str:
    return datetime.fromtimestamp(ts, ZoneInfo(tz_name)).strftime("%d.%m.%Y à %H:%M")


def _zones_du_mois(payload: dict, period_key: str) -> list[tuple]:
    out = []
    for z in payload.get("zones") or []:
        e = (z.get("periodes") or {}).get(period_key)
        if e is None:
            continue
        out.append((z, e))
    return out


def _couverture_mois(payload: dict, period_key: str, entries: list) -> tuple[str, str]:
    """Libellé du mois et phrase de période, en heure suisse. Un mois en
    cours s'arrête au dernier relevé, comme sur la page."""
    tz = payload.get("timezone") or TIMEZONE
    year, month = parse_period_key(period_key)
    label = period_label(year, month)
    start, end = period_bounds(year, month, tz)
    for p in payload.get("periodes") or []:
        if p.get("key") != period_key:
            continue
        label = p.get("label") or label
        start = p.get("start", start)
        end = p.get("end", end)
        break

    bat = ((payload.get("batiment") or {}).get("periodes") or {}).get(period_key) or {}
    generated = payload.get("generated_at")
    en_cours = bool(bat.get("en_cours")) or any(e.get("en_cours") for _, e in entries)
    if generated and end > generated:
        en_cours = True

    if not en_cours:
        return label, f"{_jour(start, tz)} → {_jour(end - 1, tz)} (heure de Suisse)"

    stamps = []
    for _, e in entries:
        for side in ("reseau", "solaire", "controle"):
            ts = (e.get(side) or {}).get("releve_fin_ts")
            if ts:
                stamps.append(ts)
    prod_ts = (bat.get("production") or {}).get("releve_fin_ts")
    if prod_ts:
        stamps.append(prod_ts)
    if stamps:
        return label, (
            f"{_jour(start, tz)} → {_jour(max(stamps), tz)} (dernier relevé, mois en cours)"
        )
    return label, f"{_jour(start, tz)} → {_jour(end - 1, tz)} (mois en cours)"


def _releves_manquants(entries: list) -> list[str]:
    """Pourquoi une ligne n'est pas facturable. Les alertes des colonnes
    d'abord ; à défaut, celles de la zone (répartition impossible, compteur
    muet). Le mois en cours est déjà dit dans le sous-titre."""
    lignes = []
    for z, e in entries:
        if etat_zone(e) == "mois en cours":
            continue
        nom = z.get("label") or z.get("zone") or ""
        cote_lignes = []
        for cote, key in (("Réseau", "reseau"), ("Solaire", "solaire")):
            for detail in (e.get(key) or {}).get("alertes") or []:
                cote_lignes.append(f"{nom} — {cote} : {detail}")
        if cote_lignes:
            lignes.extend(cote_lignes)
            continue
        if etat_zone(e) == "facturable":
            continue
        for detail in e.get("alertes") or []:
            lignes.append(f"{nom} : {detail}")
    return lignes


def _notes_classeur(payload: dict, period_key: str, manquants: list[str],
                    postes: dict | None = None) -> list[tuple[str, str]]:
    notes: list[tuple[str, str]] = [(
        "Chaque zone est scindée en kWh achetés au réseau et kWh solaires "
        "autoconsommés. La consommation totale est leur somme.",
        "note",
    )]
    if _buanderie_par_zone(postes) is not None:
        notes.append((
            "La colonne Buanderie est l'électricité de la buanderie commune "
            "attribuée au lot (CEnergieAppXX). Elle passe par le compteur des "
            "Communs : elle n'est ni dans la consommation totale du lot, ni "
            "dans les montants, ni déduite des Communs. Détail dans l'onglet Buanderie.",
            "note",
        ))
    elif postes and postes.get("buanderie_avant_mise_en_service"):
        notes.append((
            "Buanderie : pas de décompte avant sa mise en service le 17.07.2026.",
            "note",
        ))
    rep = repartition_du_mois(payload, period_key)
    if rep.get("source") == "calculee":
        notes.append((
            "Ces parts sont calculées. La consommation vient du compteur du lot, "
            "répartie selon le mix du bâtiment (production moins réinjection). "
            "Les séries Grid et Solaire du Miniserver ne sont pas utilisées.",
            "note",
        ))
    elif rep.get("source") == "loxone" and not rep.get("confiance_verifiee"):
        notes.append((
            "Ces parts sont les mesures du Miniserver. Le bilan du bâtiment "
            "ne permet pas de les revérifier.",
            "note",
        ))
    elif rep.get("source") == "loxone":
        notes.append(("Ces parts sont les mesures du Miniserver.", "note"))
    if payload.get("reserves"):
        notes.append(("Sous réserve", "head"))
        notes.extend((ligne, "note") for ligne in payload["reserves"])
    notes.append((
        "Une case vide n'est pas un zéro : le relevé manque, et le total "
        "ne compte que les relevés présents.",
        "note",
    ))
    notes.append((
        "Le total porte sur toutes les zones. Filtrer le tableau ne change pas ce total.",
        "note",
    ))
    if manquants:
        notes.append(("Relevés incomplets", "head"))
        notes.extend((ligne, "note") for ligne in manquants)
    return notes


def _style_kwh(cell, value) -> None:
    if isinstance(value, float):
        cell.number_format = _KWH_FMT
        cell.alignment = _ALIGN_RIGHT
    else:
        cell.alignment = _ALIGN_RIGHT


def _tarif_du_mois(entries: list) -> dict | None:
    """Un mois n'a qu'un tarif pour le site. On prend celui déjà résolu
    sur les zones (tarif_for au début de la période)."""
    for _, e in entries:
        tarif = e.get("tarif") or None
        if tarif and tarif.get("prix_reseau") is not None and tarif.get("prix_solaire") is not None:
            return tarif
    return None


def _formule_ttc(qte: str, prix: str, tva: str) -> str:
    """Montant TTC, vide si la quantité, le prix ou la TVA manque : un
    0 CHF ne doit pas se confondre avec une ligne sans relevé."""
    return f'=IF(OR({qte}="",{prix}="",{tva}=""),"",{qte}*{prix}*(1+{tva}/100))'


def _formule_somme(col: str, first: int, last: int) -> str:
    return f'=IF(COUNT({col}{first}:{col}{last})=0,"",SUM({col}{first}:{col}{last}))'


def _mise_en_page(ws, ncol: int, last_row: int, header_row: int,
                  site: str, label: str, footer: str) -> None:
    last_col = get_column_letter(ncol)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.75, bottom=0.6, header=0.28, footer=0.28)
    ws.print_title_rows = f"1:{header_row}"
    ws.print_area = f"A1:{last_col}{last_row}"
    ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.sheet_properties.tabColor = _NAVY
    ws.oddHeader.left.text = site or "Décompte"
    ws.oddHeader.right.text = label
    ws.oddFooter.left.text = footer
    ws.oddFooter.right.text = "Page &P / &N"
    ws.page_setup.oddEven = False


def _ecrire_titre(ws, titre: str, sous: str, ncol: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    cell = ws.cell(1, 1, titre)
    cell.font = _FONT_TITLE
    cell.alignment = _ALIGN_LEFT
    ws.row_dimensions[1].height = 26
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    cell = ws.cell(2, 1, sous)
    cell.font = _FONT_SUB
    cell.alignment = _ALIGN_WRAP
    ws.row_dimensions[2].height = 32


def _prix_jaune(cell, value, fmt: str) -> None:
    cell.value = value
    cell.fill = _FILL_EDIT
    cell.border = _THIN
    cell.font = Font(bold=True)
    cell.number_format = fmt
    cell.alignment = _ALIGN_RIGHT


def _compteurs_statistiques(zone: dict, calculee: bool) -> list[tuple[str, str]]:
    """Compteur physique de l'appartement, celui dont les statistiques
    du Miniserver sont un index. En répartition recalculée c'est le
    compteur du lot. Sinon les compteurs mesurés (réseau et solaire),
    ou le compteur unique du lot quand il n'y a pas de split."""
    sources = zone.get("sources") or {}

    def add(key, out, seen):
        src = sources.get(key) or {}
        sid = src.get("series_id")
        if not sid or sid in seen:
            return
        seen.add(sid)
        out.append((src.get("label") or key, sid))

    out, seen = [], set()
    if calculee or not (sources.get("reseau") or sources.get("solaire")):
        add("controle", out, seen)
    else:
        add("reseau", out, seen)
        add("solaire", out, seen)
    return out


def _points_horaires(conn, series_ids: list[str], start: int, end: int) -> dict[str, list[tuple]]:
    """Points Statistics du Miniserver, un par échantillon."""
    if conn is None or not series_ids:
        return {}
    marks = ",".join("?" * len(series_ids))
    rows = conn.execute(
        f"""
        SELECT series_id, ts, value FROM statistic_samples
        WHERE series_id IN ({marks}) AND ts >= ? AND ts < ?
        ORDER BY series_id, ts
        """,
        (*series_ids, start, end),
    ).fetchall()
    out: dict[str, list[tuple]] = {}
    for sid, ts, value in rows:
        out.setdefault(sid, []).append((ts, value))
    return out


def _bornes_mois(payload: dict, period_key: str) -> tuple[int, int]:
    tz = payload.get("timezone") or TIMEZONE
    year, month = parse_period_key(period_key)
    start, end = period_bounds(year, month, tz)
    for p in payload.get("periodes") or []:
        if p.get("key") == period_key:
            start = p.get("start", start)
            end = p.get("end", end)
            break
    return start, end


def _ecrire_onglet_releves(wb: Workbook, entries: list, payload: dict, period_key: str,
                           site: str, label: str, couverture: str, tz: str, conn) -> None:
    """Une ligne par point horaire des statistiques du Miniserver, pour
    le compteur de chaque appartement."""
    calculee = repartition_du_mois(payload, period_key).get("source") == "calculee"
    start, end = _bornes_mois(payload, period_key)
    compteurs = []
    for z, _e in entries:
        zone = z.get("label") or z.get("zone") or ""
        for nom, sid in _compteurs_statistiques(z, calculee):
            compteurs.append((zone, nom, sid))
    points = _points_horaires(conn, [sid for _, _, sid in compteurs], start, end)

    ws = wb.create_sheet("Relevés")
    ncol = 4
    sous = (
        f"{label}  ·  {couverture}. Une ligne par point horaire enregistré "
        "dans les statistiques du Miniserver, pour le compteur de chaque appartement."
    )
    _ecrire_titre(ws, f"Relevés de compteurs — {site}" if site else "Relevés de compteurs", sous, ncol)
    ws.row_dimensions[2].height = 36
    headers = ["Zone", "Compteur", "Date", "Index"]
    header_row = 4
    for col, nom in enumerate(headers, start=1):
        cell = ws.cell(header_row, col, nom)
        cell.font = _FONT_HEADER
        cell.fill = _FILL_HEADER
        cell.alignment = _ALIGN_CENTER
        cell.border = _THIN
    ws.row_dimensions[header_row].height = 22

    row = header_row
    for zone, nom, sid in compteurs:
        for ts, value in points.get(sid, []):
            row += 1
            ws.cell(row, 1, zone).alignment = _ALIGN_LEFT
            ws.cell(row, 2, nom).alignment = _ALIGN_LEFT
            date = ws.cell(row, 3, _horodatage(int(ts), tz))
            date.alignment = _ALIGN_CENTER
            index = ws.cell(row, 4, _index_affiche(value))
            index.number_format = _INDEX_FMT
            index.alignment = _ALIGN_RIGHT

    n_points = row - header_row
    if n_points:
        ws.auto_filter.ref = f"A{header_row}:D{row}"
    for col, width in {"A": 22, "B": 36, "C": 22, "D": 16}.items():
        ws.column_dimensions[col].width = width
    _mise_en_page(
        ws, ncol, max(row, header_row), header_row, site, label, "Relevés de compteurs",
    )


def _ecrire_notes(ws, start_row: int, notes: list[tuple[str, str]], ncol: int,
                  largeur_totale: float) -> int:
    """Notes de lecture sous un tableau, une par ligne fusionnée. Retourne
    la dernière ligne écrite."""
    car_par_ligne = max(40, int(largeur_totale * 0.73))
    row = start_row
    for text, kind in notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
        cell = ws.cell(row, 1, text)
        cell.font = _FONT_NOTE_HEAD if kind == "head" else _FONT_NOTE
        cell.alignment = _ALIGN_WRAP
        lignes_texte = max(1, (len(text) + car_par_ligne - 1) // car_par_ligne)
        ws.row_dimensions[row].height = 16 * lignes_texte + 4
        row += 1
    return row - 1


def _ecrire_onglet_tableau(wb: Workbook, nom: str, titre: str, sous: str,
                           colonnes: list[tuple[str, int, str | None]],
                           lignes: list[list], total: list | None,
                           notes: list[tuple[str, str]],
                           site: str, label: str, footer: str,
                           prix: list[tuple[str, float | None, str]] | None = None,
                           note_prix: str = "") -> None:
    """Onglet à un tableau : en-tête, une ligne par entrée, ligne de total,
    notes. `colonnes` : (titre, largeur, format) où format est un format
    numérique Excel, None pour du texte, ou "etat" pour la colonne d'état.

    `prix` : cases jaunes modifiables (libellé, valeur, format) écrites en
    B4, B5... au-dessus du tableau. Une cellule de ligne peut alors être
    une fonction `f(row)` qui rend une formule, et une cellule du total une
    fonction `f(first, last)`."""
    ws = wb.create_sheet(nom)
    ncol = len(colonnes)
    _ecrire_titre(ws, titre, sous, ncol)
    ws.row_dimensions[2].height = 36
    header_row = 4
    if prix:
        for i, (libelle, valeur, fmt) in enumerate(prix):
            r = 4 + i
            ws.cell(r, 1, libelle).font = Font(bold=True, color=_NAVY)
            _prix_jaune(ws.cell(r, 2), valeur, fmt)
            ws.row_dimensions[r].height = 20
        last_prix = 4 + len(prix) - 1
        ws.merge_cells(start_row=4, start_column=3, end_row=last_prix, end_column=ncol)
        cell = ws.cell(4, 3, note_prix)
        cell.font = _FONT_SUB
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        header_row = last_prix + 2
    for c, (nom_col, _w, _f) in enumerate(colonnes, start=1):
        cell = ws.cell(header_row, c, nom_col)
        cell.font = _FONT_HEADER
        cell.fill = _FILL_HEADER
        cell.alignment = _ALIGN_CENTER
        cell.border = _THIN
    ws.row_dimensions[header_row].height = 32

    i_etat = next((i for i, (_n, _w, f) in enumerate(colonnes) if f == "etat"), None)
    row = header_row
    for i, valeurs in enumerate(lignes):
        row += 1
        etat = valeurs[i_etat] if i_etat is not None else "facturable"
        fill = _FILL_AMBER if etat != "facturable" else (_FILL_STRIPE if i % 2 else None)
        for c, (value, (_n, _w, fmt)) in enumerate(zip(valeurs, colonnes), start=1):
            formule = callable(value)
            cell = ws.cell(row, c, value(row) if formule else value)
            cell.border = _THIN
            if fill is not None:
                cell.fill = fill
            if fmt == "etat":
                cell.font = _FONT_ETAT.get(value, Font())
                cell.alignment = _ALIGN_CENTER
            elif fmt:
                if formule or isinstance(value, (int, float)):
                    cell.number_format = fmt
                cell.alignment = _ALIGN_RIGHT
            else:
                cell.alignment = _ALIGN_LEFT
        ws.row_dimensions[row].height = 18
    if lignes:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncol)}{row}"

    if total is not None and lignes:
        first_data, last_data = header_row + 1, row
        row += 1
        for c, (value, (_n, _w, fmt)) in enumerate(zip(total, colonnes), start=1):
            formule = callable(value)
            cell = ws.cell(row, c, value(first_data, last_data) if formule else value)
            cell.font = _FONT_TOTAL
            cell.fill = _FILL_TOTAL
            cell.border = _TOTAL_BORDER
            if fmt and fmt != "etat":
                if formule or isinstance(value, (int, float)):
                    cell.number_format = fmt
                cell.alignment = _ALIGN_RIGHT
            else:
                cell.alignment = _ALIGN_LEFT
        ws.row_dimensions[row].height = 22

    for c, (_n, width, _f) in enumerate(colonnes, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    last_row = _ecrire_notes(ws, row + 2, notes, ncol, sum(w for _n, w, _f in colonnes))
    _mise_en_page(ws, ncol, max(last_row, header_row), header_row, site, label, footer)


def _notes_manquants(lignes: list[dict], nom) -> list[tuple[str, str]]:
    """Relevés qui empêchent une ligne d'être facturable. Le mois en cours
    est déjà dit dans le sous-titre."""
    detail = []
    for ln in lignes:
        if ln["etat"] != "données incomplètes":
            continue
        for a in ln["alertes"] or ["relevé manquant"]:
            detail.append((f"{nom(ln)} : {a}", "note"))
    return [("Relevés incomplets", "head"), *detail] if detail else []


def _index_identiques(lignes: list[dict]) -> list[tuple[str, str]]:
    """Compteurs de zones différentes aux mêmes index de début et de fin :
    à Sequoia, l'eau chaude de l'App 41 a l'historique exact de celle de
    l'App 35 (un seul compteur lu deux fois dans Loxone Config)."""
    groupes: dict[tuple, list[dict]] = {}
    for ln in lignes:
        if ln["debut"] and ln["fin"] is not None:
            groupes.setdefault((ln["debut"], ln["fin"]), []).append(ln)
    notes = []
    for (debut, fin), group in groupes.items():
        zones = list(dict.fromkeys(ln["zone"] for ln in group))
        if len(zones) < 2:
            continue
        notes.append((
            f"{' et '.join(zones)} : mêmes index en début et en fin de mois "
            f"({_fr(debut, _INDEX_DECIMALES)} → {_fr(fin, _INDEX_DECIMALES)}). "
            "Probablement un seul compteur lu deux fois : à vérifier dans Loxone Config.",
            "note",
        ))
    return [("Points d'attention", "head"), *notes] if notes else []


def _ecrire_onglets_postes(wb: Workbook, postes: dict, site: str, label: str,
                           couverture: str, tva_defaut: float | None = None) -> None:
    """Onglets Chauffage, Eau chaude et Buanderie. Un onglet n'est créé que
    si le site a des compteurs de ce poste.

    Chauffage et eau chaude ont leur prix HT et leur TVA en cases jaunes
    (pas de tarif en base pour ces postes : le prix part vide, la TVA
    reprend celle du tarif électrique s'il y en a un)."""
    explications = {
        "chauffage": "Énergie de chauffage = index de fin − index de début du "
                     "compteur de chaleur. Le volume circulé n'est pas repris.",
        "eau_chaude": "Volume d'eau chaude sanitaire = index de fin − index de "
                      "début. L'eau froide n'est pas reprise.",
    }
    for role, (titre, grandeur, unite, dec) in _POSTE_TABLES.items():
        lignes, total = lignes_compteurs(postes, role)
        if not lignes:
            continue
        qte_fmt = _M3_FMT if unite == "m³" else _KWH_FMT
        colonnes = [
            ("Zone", 22, None),
            ("Compteur", 46, None),
            (f"Index début ({unite})", 18, _INDEX_FMT),
            (f"Index fin ({unite})", 18, _INDEX_FMT),
            (f"{grandeur} ({unite})", 18, qte_fmt),
            ("Montant TTC (CHF)", 18, _CHF_FMT),
            ("État", 22, "etat"),
        ]
        # Prix en B4, TVA en B5 ; quantité en E, montant en F.
        def montant(row):
            return _formule_ttc(f"E{row}", "$B$4", "$B$5")

        def somme_montants(first, last):
            return _formule_somme("F", first, last)

        rows = [[ln["zone"], ln["compteur"], ln["debut"], ln["fin"], ln["quantite"],
                 montant, ln["etat"]]
                for ln in lignes]
        notes = [(explications[role], "note")]
        if len({ln["zone"] for ln in lignes}) < len(lignes):
            notes.append((
                "Une zone peut avoir plusieurs compteurs : une ligne par compteur, "
                "le total les additionne tous.", "note",
            ))
        notes.append((
            "Une case vide n'est pas un zéro : le relevé manque, et le total "
            "ne compte que les relevés présents.", "note",
        ))
        notes += _index_identiques(lignes)
        notes += _notes_manquants(lignes, lambda ln: f"{ln['zone']} — {ln['compteur']}")
        _ecrire_onglet_tableau(
            wb, titre, f"{titre} — {site}" if site else titre,
            f"{label}  ·  {couverture}", colonnes, rows,
            ["Total immeuble", "", None, None, total, somme_montants, ""], notes,
            site, label, f"{titre} · {unite} et CHF",
            prix=[
                (f"Prix {titre.lower()} HT (CHF/{unite})", None, _PRIX_FMT),
                ("TVA (%)", tva_defaut, _TVA_FMT),
            ],
            note_prix=(
                f"Montant TTC = {grandeur.lower()} × prix HT × (1 + TVA). Les cases "
                "jaunes sont modifiables : les montants se recalculent. Une quantité, "
                "un prix ou une TVA vide ne donne pas de montant."
            ),
        )

    lignes, totaux = lignes_buanderie(postes)
    if not lignes:
        return
    colonnes = [
        ("Zone", 22, None),
        ("Cycles", 14, _CYCLES_FMT),
        ("Énergie (kWh)", 18, _KWH_FMT),
        ("État", 22, "etat"),
    ]
    rows = [[ln["zone"], ln["cycles"], ln["energie"], ln["etat"]] for ln in lignes]
    notes = [
        ("Énergie (CEnergieAppXX) et nombre de cycles (CNrMachineAppXX) de la "
         "buanderie commune, attribués à chaque appartement par badge.", "note"),
        ("Cette énergie passe par le compteur des Communs : elle n'est pas dans "
         "la consommation électrique du lot, et n'est pas déduite des Communs.", "note"),
    ]
    if postes.get("buanderie_depuis"):
        notes.append((
            "Décompte à partir du 17.07.2026, mise en service de la buanderie. "
            "Avant, l'énergie s'incrémentait sur tous les lots sans cycle.", "note",
        ))
    notes += _notes_manquants(lignes, lambda ln: ln["zone"])
    _ecrire_onglet_tableau(
        wb, "Buanderie", f"Buanderie — {site}" if site else "Buanderie",
        f"{label}  ·  {couverture}", colonnes, rows,
        ["Total immeuble", totaux["cycles"], totaux["energie"], ""], notes,
        site, label, "Buanderie · kWh et cycles",
    )


_FRAIS_COLONNES = [
    # (titre, largeur, saisie par l'utilisateur)
    ("Zone", 28, False),
    ("Location des compteurs (CHF)", 20, True),
    ("Frais exceptionnels (CHF)", 20, True),
    ("Autres frais (CHF)", 20, True),
    ("Total frais (CHF)", 18, False),
    ("Remarque", 46, True),
]


def _ecrire_onglet_frais(wb: Workbook, zones: list[tuple[str, int]], site: str,
                         label: str, couverture: str) -> None:
    """Tableau vide, une ligne par zone de l'onglet Électricité, pour
    reporter à la main les frais hors relevés (location des compteurs,
    frais exceptionnels). `zones` : (nom, ligne dans Électricité) ; le nom
    de la zone est un lien vers sa ligne."""
    ws = wb.create_sheet("Frais supplémentaires")
    ncol = len(_FRAIS_COLONNES)
    _ecrire_titre(
        ws, f"Frais supplémentaires — {site}" if site else "Frais supplémentaires",
        f"{label}  ·  {couverture}", ncol,
    )
    ws.row_dimensions[2].height = 22
    header_row = 4
    for c, (nom, _w, _s) in enumerate(_FRAIS_COLONNES, start=1):
        cell = ws.cell(header_row, c, nom)
        cell.font = _FONT_HEADER
        cell.fill = _FILL_HEADER
        cell.alignment = _ALIGN_CENTER
        cell.border = _THIN
    ws.row_dimensions[header_row].height = 32

    first = header_row + 1
    for i, (zone, ligne_elec) in enumerate(zones):
        row = first + i
        for c, (_n, _w, saisie) in enumerate(_FRAIS_COLONNES, start=1):
            cell = ws.cell(row, c)
            cell.border = _THIN
            if saisie:
                cell.fill = _FILL_EDIT
            if c == 1:
                cell.value = zone
                cell.hyperlink = Hyperlink(ref=cell.coordinate,
                                           location=f"'Électricité'!A{ligne_elec}")
                cell.font = Font(color="0563C1", underline="single")
                cell.alignment = _ALIGN_LEFT
            elif c == ncol:
                cell.alignment = _ALIGN_WRAP
            else:
                cell.number_format = _CHF_FMT
                cell.alignment = _ALIGN_RIGHT
        ws.cell(row, 5).value = f'=IF(COUNT(B{row}:D{row})=0,"",SUM(B{row}:D{row}))'
        ws.cell(row, 5).font = Font(bold=True)
        ws.row_dimensions[row].height = 18
    last = first + len(zones) - 1

    total_row = last + 1
    for c in range(1, ncol + 1):
        cell = ws.cell(total_row, c)
        cell.font = _FONT_TOTAL
        cell.fill = _FILL_TOTAL
        cell.border = _TOTAL_BORDER
        if 2 <= c <= 5:
            cell.value = _formule_somme(get_column_letter(c), first, last)
            cell.number_format = _CHF_FMT
            cell.alignment = _ALIGN_RIGHT
    ws.cell(total_row, 1, "Total immeuble").alignment = _ALIGN_LEFT
    ws.row_dimensions[total_row].height = 22

    for c, (_n, width, _s) in enumerate(_FRAIS_COLONNES, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    notes = [
        ("Tableau à remplir : les cases jaunes sont vides à la génération. "
         "Reporter les montants à facturer en CHF, TVA comprise.", "note"),
        ("Le nom de chaque zone renvoie à sa ligne dans l'onglet Électricité. "
         "Le total de la ligne et celui de l'immeuble se calculent seuls.", "note"),
    ]
    last_row = _ecrire_notes(ws, total_row + 2, notes, ncol,
                             sum(w for _n, w, _s in _FRAIS_COLONNES))
    _mise_en_page(ws, ncol, last_row, header_row, site, label, "Frais supplémentaires · CHF")


def build_decompte_workbook(payload: dict, period_key: str, site: str, conn=None) -> Workbook:
    """Classeur du mois choisi, ouvrable tel quel.

    Onglets : Électricité (kWh, prix et montants, colonne buanderie si le
    mois en a une), puis Chauffage, Eau chaude et Buanderie quand le site
    a ces compteurs (`payload["postes"]`), Frais supplémentaires (tableau
    vide à remplir) et enfin les Relevés horaires.

    Les chiffres sont ceux de `export_lignes` (donc de la page), arrondis
    au centième de kWh pour que la somme des lignes égale le total affiché.
    Lève KeyError si le mois n'a aucune zone.
    """
    entries = _zones_du_mois(payload, period_key)
    if not entries:
        raise KeyError(period_key)

    tz = payload.get("timezone") or TIMEZONE
    label, couverture = _couverture_mois(payload, period_key, entries)
    generated = payload.get("generated_at")
    manquants = _releves_manquants(entries)
    postes = _postes_du_mois(payload, period_key)
    buanderie = _buanderie_par_zone(postes)

    lignes = []
    for z, e in entries:
        lignes.append({
            "zone": z.get("label") or z.get("zone") or "",
            "reseau": _kwh_affiche((e.get("reseau") or {}).get("kwh")),
            "solaire": _kwh_affiche((e.get("solaire") or {}).get("kwh")),
            "total": _kwh_affiche(e.get("total")),
            "buanderie": _kwh_affiche((buanderie or {}).get(z.get("zone"))),
            "etat": etat_zone(e),
        })
    totaux = {k: _somme_affichee([ln[k] for ln in lignes]) for k in _ELEC_KWH}
    taux = None
    if totaux["total"]:
        taux = (totaux["solaire"] or 0.0) / totaux["total"]

    cles = ["zone", "total", "solaire", "reseau"]
    if buanderie is not None:
        cles.append("buanderie")
    cles += [*_ELEC_CHF, "etat"]
    col = {k: i for i, k in enumerate(cles, start=1)}
    ncol = len(cles)

    def ref(key: str, row: int) -> str:
        return f"{get_column_letter(col[key])}{row}"

    wb = Workbook()
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    ws = wb.active
    ws.title = "Électricité"

    sous = label + "  ·  " + couverture
    if generated:
        sous += f"  ·  généré le {_horodatage(int(generated), tz)}"
    _ecrire_titre(
        ws,
        f"Décompte de charges — {site}" if site else "Décompte de charges",
        sous, ncol,
    )
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 8

    tarif = _tarif_du_mois(entries)
    prix_reseau = tarif.get("prix_reseau") if tarif else None
    prix_solaire = tarif.get("prix_solaire") if tarif else None
    taux_tva = tarif.get("taux_tva") if tarif else None
    ws.cell(_PRIX_RESEAU_ROW, 1, "Prix réseau HT (CHF/kWh)").font = Font(bold=True, color=_NAVY)
    ws.cell(_PRIX_SOLAIRE_ROW, 1, "Prix solaire HT (CHF/kWh)").font = Font(bold=True, color=_NAVY)
    ws.cell(_TVA_ROW, 1, "TVA (%)").font = Font(bold=True, color=_NAVY)
    _prix_jaune(ws.cell(_PRIX_RESEAU_ROW, 2), prix_reseau, _PRIX_FMT)
    _prix_jaune(ws.cell(_PRIX_SOLAIRE_ROW, 2), prix_solaire, _PRIX_FMT)
    _prix_jaune(ws.cell(_TVA_ROW, 2), taux_tva, _TVA_FMT)
    for r in (_PRIX_RESEAU_ROW, _PRIX_SOLAIRE_ROW, _TVA_ROW):
        ws.row_dimensions[r].height = 20
    vf = str((tarif or {}).get("valid_from") or "")
    if len(vf) >= 10 and vf[4] == "-" and vf[7] == "-":
        origine = f"Tarif en vigueur depuis le {vf[8:10]}.{vf[5:7]}.{vf[:4]}."
    elif tarif:
        origine = "Tarif enregistré pour ce site."
    else:
        origine = "Aucun tarif enregistré pour ce mois."
    ws.merge_cells(start_row=_PRIX_RESEAU_ROW, start_column=3, end_row=_TVA_ROW, end_column=ncol)
    note_prix = ws.cell(
        _PRIX_RESEAU_ROW, 3,
        f"{origine} Les montants sont TTC : kWh × prix HT × (1 + TVA). "
        "Les cases jaunes sont modifiables : les montants se recalculent. "
        "Une quantité vide ne donne pas de montant.",
    )
    note_prix.font = _FONT_SUB
    note_prix.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ref_prix_reseau = f"$B${_PRIX_RESEAU_ROW}"
    ref_prix_solaire = f"$B${_PRIX_SOLAIRE_ROW}"
    ref_tva = f"$B${_TVA_ROW}"

    for key in cles:
        cell = ws.cell(_XLSX_HEADER_ROW, col[key], _ELEC_COLONNES[key][0])
        cell.font = _FONT_HEADER
        cell.fill = _FILL_HEADER
        cell.alignment = _ALIGN_CENTER
        cell.border = _THIN
    ws.row_dimensions[_XLSX_HEADER_ROW].height = 32

    first_data = _XLSX_HEADER_ROW + 1
    for i, ln in enumerate(lignes):
        row = first_data + i
        incomplet = ln["etat"] != "facturable"
        fill = _FILL_AMBER if incomplet else (_FILL_STRIPE if i % 2 else None)
        chf_reseau, chf_solaire = ref("chf_reseau", row), ref("chf_solaire", row)
        valeurs = {
            "zone": ln["zone"],
            "total": ln["total"],
            "solaire": ln["solaire"],
            "reseau": ln["reseau"],
            "buanderie": ln["buanderie"],
            "chf_reseau": _formule_ttc(ref("reseau", row), ref_prix_reseau, ref_tva),
            "chf_solaire": _formule_ttc(ref("solaire", row), ref_prix_solaire, ref_tva),
            "ttc": f'=IF(OR({chf_reseau}="",{chf_solaire}=""),"",{chf_reseau}+{chf_solaire})',
            "etat": ln["etat"],
        }
        for key in cles:
            cell = ws.cell(row, col[key], valeurs[key])
            cell.border = _THIN
            if fill is not None:
                cell.fill = fill
            if key in _ELEC_KWH:
                _style_kwh(cell, valeurs[key])
                if key == "total" and valeurs[key] is not None:
                    cell.font = Font(bold=True)
            elif key in _ELEC_CHF:
                cell.number_format = _CHF_FMT
                cell.alignment = _ALIGN_RIGHT
                if key == "ttc":
                    cell.font = Font(bold=True)
            elif key == "etat":
                cell.font = _FONT_ETAT.get(ln["etat"], Font())
                cell.alignment = _ALIGN_CENTER
            else:
                cell.alignment = _ALIGN_LEFT
        ws.row_dimensions[row].height = 18

    last_data = first_data + len(lignes) - 1
    total_row = last_data + 1
    for key in cles:
        cell = ws.cell(total_row, col[key], None)
        cell.font = _FONT_TOTAL
        cell.fill = _FILL_TOTAL
        cell.border = _TOTAL_BORDER
        if key == "zone":
            cell.value = "Total immeuble"
            cell.alignment = _ALIGN_LEFT
        elif key in _ELEC_KWH:
            cell.value = totaux[key]
            _style_kwh(cell, totaux[key])
        elif key in _ELEC_CHF:
            letter = get_column_letter(col[key])
            cell.value = _formule_somme(letter, first_data, last_data)
            cell.number_format = _CHF_FMT
            cell.alignment = _ALIGN_RIGHT
        else:
            cell.alignment = _ALIGN_CENTER
    ws.row_dimensions[total_row].height = 22

    ws.auto_filter.ref = f"A{_XLSX_HEADER_ROW}:{get_column_letter(ncol)}{last_data}"

    rate_row = total_row + 2
    if taux is not None:
        phrase_taux = (
            f"Taux d'autoproduction : {taux * 100:.0f} % — part de la "
            "consommation couverte par le solaire autoconsommé."
        )
    else:
        phrase_taux = (
            "Taux d'autoproduction : pas calculé, la consommation du mois "
            "n'est pas connue."
        )
    ws.merge_cells(start_row=rate_row, start_column=1, end_row=rate_row, end_column=ncol)
    expl = ws.cell(rate_row, 1, phrase_taux)
    expl.font = Font(bold=True, size=12, color=_NAVY)
    expl.alignment = _ALIGN_WRAP
    ws.row_dimensions[rate_row].height = 22

    largeur_zone = max((len(ln["zone"]) for ln in lignes), default=12)
    largeurs = {k: _ELEC_COLONNES[k][1] for k in cles}
    largeurs["zone"] = min(42, max(28, largeur_zone + 3))
    for key, width in largeurs.items():
        ws.column_dimensions[get_column_letter(col[key])].width = width

    last_row = _ecrire_notes(
        ws, rate_row + 2, _notes_classeur(payload, period_key, manquants, postes), ncol,
        sum(largeurs.values()),
    )
    _mise_en_page(ws, ncol, last_row, _XLSX_HEADER_ROW, site, label, "Électricité · kWh et CHF")

    if postes:
        _ecrire_onglets_postes(wb, postes, site, label, couverture, taux_tva)
    _ecrire_onglet_frais(
        wb, [(ln["zone"], first_data + i) for i, ln in enumerate(lignes)],
        site, label, couverture,
    )
    _ecrire_onglet_releves(
        wb, entries, payload, period_key, site, label, couverture, tz, conn,
    )

    wb.properties.title = f"Décompte {site} {label}".strip()
    wb.properties.creator = "Loxone Collector"
    return wb
