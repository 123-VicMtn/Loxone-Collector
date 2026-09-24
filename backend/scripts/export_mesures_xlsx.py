#!/usr/bin/env python3
"""
Génère un classeur Excel de mesures (kWh) à partir des données réelles du
collecteur, sur le même principe que docs/Mesure PPE Horizon T2 2026.xlsx
(un exemple fourni par l'utilisateur pour un autre immeuble) : un onglet
"brut" du compteur d'alimentation, un onglet de synthèse mensuelle par zone,
et un onglet de détail horaire par zone.

Différence volontaire avec l'exemple : celui-ci dérivait la part PV/Réseau
de chaque appartement par une répartition proportionnelle (un seul compteur
Total par appartement, splitté a posteriori selon le mix du bâtiment à cet
instant). Ici chaque zone a son propre compteur EFM Grid et Solaire (posés
en octobre 2025) -- la part réseau et la part solaire autoconsommée sont
DIRECTEMENT mesurées, pas déduites. Voir billing.py pour la validation
empirique de ce que mesure chaque compteur.

Usage :
    python3 scripts/export_mesures_xlsx.py <config.yaml> [--miniserver NOM]
        [--out chemin.xlsx] [--from-mois AAAA-MM] [--to-mois AAAA-MM]

--miniserver : le site à exporter (nom tel que dans config.yaml). Requis dès
que la config en liste plusieurs -- une config qui suit deux immeubles (ex:
config.external.yaml : MS-Arlopi + MS-PPE-Horizon) mélangerait sinon les
zones des deux bâtiments dans un seul classeur. Optionnel si un seul
miniserver est configuré.

--from-mois / --to-mois restreignent les 3 onglets à une plage de mois
calendaires (bornes incluses, ex: --from-mois 2026-07 --to-mois 2026-08 pour
juillet et août 2026 seulement). Par défaut : tout l'historique disponible.

Lire une base SQLite en écriture (WAL) échoue via le pont device-bridge
d'une session Claude distante -- lance ce script depuis un Terminal natif
sur cette machine si besoin (voir CLAUDE.md, "Piège d'environnement").
"""
from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path
from zoneinfo import ZoneInfo

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import billing
import db
from config import load_config

TZ = ZoneInfo(billing.TIMEZONE)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GROUP_FONT = Font(bold=True, color="FFFFFF")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
NOTE_FONT = Font(italic=True, color="808080")
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center")
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
WRAP_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)


def add_title(ws, text: str, last_col: int, row: int = 1) -> None:
    """Grand titre en haut d'un onglet, fusionné sur toute la largeur du
    tableau -- pense pour un propriétaire qui ouvre le fichier sans contexte."""
    ws.cell(row=row, column=1, value=text).font = TITLE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(last_col, 1))


# --------------------------------------------------------------------------
# Lecture des relevés cumulatifs, à la résolution horaire
# --------------------------------------------------------------------------

def hourly_readings(conn, series_id: str, start_ts: int, end_ts: int) -> dict[int, float]:
    """Pour une série cumulative ('total'/'totalNeg'), le relevé le plus
    représentatif de chaque heure : le maximum observé dans l'heure (données
    brutes récentes non archivées, ou max_value archivé par
    downsample_and_prune/backfill_statistics). Un compteur cumulatif est un
    index strictement croissant, donc son maximum sur l'heure est aussi son
    relevé de fin d'heure -- même principe que db.query_daily_last, appliqué
    à l'heure plutôt qu'au jour."""
    cur = conn.execute(
        """
        SELECT hour, MAX(value) FROM (
            SELECT CAST(ts / 3600 AS INTEGER) * 3600 AS hour, value
              FROM readings
             WHERE series_id = ? AND ts BETWEEN ? AND ? AND value IS NOT NULL
            UNION ALL
            SELECT ts AS hour, max_value AS value
              FROM readings_hourly
             WHERE series_id = ? AND ts BETWEEN ? AND ?
        )
        GROUP BY hour ORDER BY hour
        """,
        (series_id, start_ts, end_ts, series_id, start_ts, end_ts),
    )
    return dict(cur.fetchall())


def series_range(conn, series_id: str) -> tuple[int, int] | None:
    row = conn.execute(
        "SELECT MIN(ts), MAX(ts) FROM readings_hourly WHERE series_id = ?", (series_id,)
    ).fetchone()
    firsts = [row[0]] if row and row[0] is not None else []
    lasts = [row[1]] if row and row[1] is not None else []
    row2 = conn.execute(
        "SELECT MAX(ts) FROM readings WHERE series_id = ? AND value IS NOT NULL", (series_id,)
    ).fetchone()
    if row2 and row2[0] is not None:
        lasts.append(row2[0])
    if not firsts:
        return None
    return min(firsts), max(lasts)


def fmt_datetime(ts: int):
    """Date + heure de la mesure, en une seule valeur (heure locale
    Europe/Zurich) -- une cellule Excel unique, sans ambiguïté sur l'heure
    à laquelle le relevé a été fait."""
    return __import__("datetime").datetime.fromtimestamp(ts, TZ).replace(tzinfo=None)


# --------------------------------------------------------------------------
# Onglet "Lisez-moi" -- légende en langage clair pour le propriétaire
# --------------------------------------------------------------------------

def build_legende(wb: Workbook, site_title: str, generated_at: int, periods: list[dict], zones: list[dict]) -> None:
    ws = wb.create_sheet("Lisez-moi", 0)
    add_title(ws, f"Mesures énergétiques -- {site_title}", 2)

    date_str = __import__("datetime").datetime.fromtimestamp(generated_at, TZ).strftime("%d.%m.%Y à %H:%M")
    periode_str = f"{periods[0]['label']} -- {periods[-1]['label']}" if periods else ""
    zones_str = ", ".join(z["label"] for z in zones)

    rows = [
        ("", ""),
        ("Généré le", f"{date_str} (heure de Suisse)"),
        ("Période couverte", periode_str),
        ("Zones incluses", zones_str),
        ("", ""),
        ("Comment lire ce classeur", ""),
        (
            "Réseau (acheté)",
            "Électricité achetée au fournisseur, consommée par la zone.",
        ),
        (
            "Solaire (autoconsommé)",
            "Électricité produite par les panneaux solaires du bâtiment et "
            "consommée directement par la zone (pas de détour par le réseau).",
        ),
        (
            "Consommation totale",
            "Réseau + Solaire -- tout ce que la zone a réellement consommé.",
        ),
        (
            "Taux d'autoproduction",
            "Part de la consommation d'une zone couverte par le solaire. "
            "Monte l'été, descend l'hiver.",
        ),
        (
            "Injection réseau",
            "Production solaire non consommée sur place, revendue/réinjectée "
            "dans le réseau électrique.",
        ),
        ("", ""),
        ("Onglet « Résumé »", "Un chiffre par mois et par zone -- la vue d'ensemble."),
        (
            "Onglet « Compteur Alimentation »",
            "Relevé heure par heure du compteur général du bâtiment (production "
            "solaire, achat au réseau) -- sert de contrôle, pas de base pour la "
            "facturation par zone.",
        ),
        (
            "Onglet « Détail par zone »",
            "Relevé heure par heure de chaque zone -- pour vérifier un pic de "
            "consommation ou une journée précise.",
        ),
        ("", ""),
        (
            "Mois « en cours »",
            "Affiché en gris clair dans le Résumé : le mois n'est pas terminé, "
            "les chiffres vont encore bouger.",
        ),
    ]
    r = 3
    for label, text in rows:
        if label and not text:
            c = ws.cell(row=r, column=1, value=label)
            c.font = BOLD
        else:
            ws.cell(row=r, column=1, value=label).font = BOLD
            cell = ws.cell(row=r, column=2, value=text)
            cell.alignment = WRAP_LEFT
        r += 1

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90


# --------------------------------------------------------------------------
# Onglet "Compteur Alimentation" -- compteurs de bâtiment, relevé horaire
# --------------------------------------------------------------------------

def build_compteur_alimentation(wb: Workbook, conn, batiment_src: dict, now_ts: int,
                                 window: tuple[int, int] | None = None) -> None:
    ws = wb.create_sheet("Compteur Alimentation")

    cols = [
        ("production", "Production solaire"),
        ("reseau_import", "Achat au réseau"),
    ]
    present = [(key, label) for key, label in cols if batiment_src.get(key)]
    # "Réseau -- Export" (totalNeg) est exclu volontairement : la série entière
    # vaut une constante ~4 294 967.295 kWh (sentinelle de dépassement Uint32),
    # jamais une vraie mesure -- vérifié empiriquement (voir series_range /
    # hourly_readings sur cette série). billing.py ne l'utilise déjà pour rien
    # (l'injection se déduit de production - autoconsommation).

    last_col = 1 + len(present) * 2
    add_title(ws, "Compteur général du bâtiment -- relevé heure par heure", last_col)
    ws["A2"] = (
        "Compteur de contrôle au raccordement de l'onduleur (ne couvre pas "
        "exactement les 6 zones -- voir l'onglet « Détail par zone » pour la "
        "consommation réelle de chaque zone)."
    )
    ws["A2"].font = NOTE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)

    header_row = 4
    ws.cell(row=header_row, column=1, value="Date et heure de la mesure")
    col = 2
    delta_cols = []
    for key, label in present:
        c = ws.cell(row=header_row, column=col, value=f"{label} -- relevé cumulé (kWh)")
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        col += 1
    for key, label in present:
        c = ws.cell(row=header_row, column=col, value=f"{label} -- énergie sur l'heure (kWh)")
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        delta_cols.append(col)
        col += 1
    for c in range(1, col):
        cell = ws.cell(row=header_row, column=c)
        if cell.font != HEADER_FONT:
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL

    series_data = {}
    display_hours = set()
    for key, _label in present:
        sid = batiment_src[key]["series_id"]
        rng = series_range(conn, sid)
        if not rng:
            series_data[key] = {}
            continue
        q_start = max(rng[0], window[0] - 3600) if window else rng[0]
        q_end = min(rng[1], now_ts, window[1]) if window else min(rng[1], now_ts)
        data = hourly_readings(conn, sid, q_start, q_end)
        series_data[key] = data
        display_start = max(rng[0], window[0]) if window else rng[0]
        display_hours.update(h for h in data.keys() if h >= display_start)

    r = header_row + 1
    prev = {key: None for key, _ in present}
    # Amorce le delta de la première ligne affichée avec le dernier relevé
    # juste avant la fenêtre, s'il existe (sinon Δ vide sur la 1re ligne).
    for key, _label in present:
        before = [h for h in series_data.get(key, {}) if h < (min(display_hours) if display_hours else 0)]
        if before:
            prev[key] = series_data[key][max(before)]
    for hour in sorted(display_hours):
        dcell = ws.cell(row=r, column=1, value=fmt_datetime(hour))
        dcell.number_format = "dd.mm.yyyy hh:mm"
        col = 2
        for key, _label in present:
            val = series_data[key].get(hour)
            ws.cell(row=r, column=col, value=round(val, 3) if val is not None else None)
            col += 1
        for key, _label in present:
            val = series_data[key].get(hour)
            p = prev[key]
            delta = round(val - p, 3) if (val is not None and p is not None) else None
            ws.cell(row=r, column=col, value=delta)
            col += 1
            if val is not None:
                prev[key] = val
        r += 1

    ws.freeze_panes = f"B{header_row + 1}"
    ws.column_dimensions["A"].width = 20
    for c in range(2, col):
        ws.column_dimensions[get_column_letter(c)].width = 26


# --------------------------------------------------------------------------
# Onglet "Résumé" -- synthèse mensuelle par zone (moteur : billing.py)
# --------------------------------------------------------------------------

def build_resume(wb: Workbook, site_title: str, decompte: dict) -> None:
    ws = wb.create_sheet("Résumé", 0)
    periods = decompte["periodes"]
    zones = decompte["zones"]
    batiment = decompte["batiment"]

    total_col_preview = 3 + len(periods)
    add_title(ws, f"Résumé mensuel de consommation -- {site_title}", total_col_preview)
    ws["A2"] = (
        f"Valeurs en kWh, généré le "
        f"{__import__('datetime').datetime.fromtimestamp(decompte['generated_at'], TZ).strftime('%d.%m.%Y à %H:%M')} "
        f"(heure de Suisse). Détail des termes dans l'onglet « Lisez-moi »."
    )
    ws["A2"].font = NOTE_FONT
    ws["A3"] = (
        "Consommation totale d'une zone = Réseau (acheté) + Solaire (autoconsommé). "
        "Un mois affiché en gris est encore en cours et n'est pas définitif."
    )
    ws["A3"].font = NOTE_FONT

    header_row = 4
    ws.cell(row=header_row, column=1, value="Zone")
    ws.cell(row=header_row, column=2, value="Grandeur")
    for i, p in enumerate(periods):
        label = p["label"]
        ws.cell(row=header_row, column=3 + i, value=label)
    total_col = 3 + len(periods)
    ws.cell(row=header_row, column=total_col, value="Total période")
    for c in range(1, total_col + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = CENTER

    r = header_row + 1
    ROWS = [
        ("total", "Consommation totale"),
        ("solaire", "  dont Solaire autoconsommé"),
        ("reseau", "  dont Réseau (achat)"),
        ("taux_autoproduction", "  Taux d'autoproduction (%)"),
    ]
    for z in zones:
        first_row = r
        for key, label in ROWS:
            ws.cell(row=r, column=1, value=z["label"] if key == "total" else "")
            ws.cell(row=r, column=2, value=label)
            total = 0.0
            total_ok = True
            for i, p in enumerate(periods):
                e = z["periodes"][p["key"]]
                if key == "total":
                    val = e["total"]
                elif key == "taux_autoproduction":
                    val = e["taux_autoproduction"]
                else:
                    val = e[key]["kwh"]
                cell = ws.cell(row=r, column=3 + i, value=round(val, 2) if val is not None else None)
                if key != "taux_autoproduction":
                    if val is None:
                        total_ok = False
                    else:
                        total += val
                if e.get("en_cours"):
                    cell.font = NOTE_FONT
            if key != "taux_autoproduction":
                ws.cell(row=r, column=total_col, value=round(total, 2) if total_ok else None)
            r += 1
        ws.cell(row=first_row, column=1).font = BOLD
        r += 1  # ligne vide entre zones

    # Synthèse bâtiment
    r += 1
    ws.cell(row=r, column=1, value="Bâtiment").font = BOLD
    ws.cell(row=r, column=1).fill = SUBHEADER_FILL
    ws.cell(row=r, column=2).fill = SUBHEADER_FILL
    for i in range(len(periods) + 1):
        ws.cell(row=r, column=3 + i).fill = SUBHEADER_FILL
    r += 1
    BATIMENT_ROWS = [
        ("production", "Production solaire totale"),
        ("autoconsommation", "  dont autoconsommée (somme zones)"),
        ("achat_reseau", "  dont achetée au réseau (somme zones)"),
        ("injection", "Injection réseau (production - autoconso.)"),
        ("taux_autoproduction", "Taux d'autoproduction bâtiment (%)"),
        ("taux_autoconsommation", "Taux d'autoconsommation bâtiment (%)"),
    ]
    for key, label in BATIMENT_ROWS:
        ws.cell(row=r, column=2, value=label)
        total = 0.0
        total_ok = True
        for i, p in enumerate(periods):
            e = batiment["periodes"][p["key"]]
            if key == "production":
                val = e["production"]["kwh"]
            else:
                val = e.get(key)
            cell = ws.cell(row=r, column=3 + i, value=round(val, 2) if val is not None else None)
            if key in ("taux_autoproduction", "taux_autoconsommation"):
                continue
            if val is None:
                total_ok = False
            else:
                total += val
        if key not in ("taux_autoproduction", "taux_autoconsommation"):
            ws.cell(row=r, column=total_col, value=round(total, 2) if total_ok else None)
        r += 1

    ws.freeze_panes = "C5"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 34
    for i in range(len(periods) + 1):
        ws.column_dimensions[get_column_letter(3 + i)].width = 14


# --------------------------------------------------------------------------
# Onglet "Détail par zone" -- relevés horaires Réseau/Solaire par zone
# --------------------------------------------------------------------------

def build_detail_zones(wb: Workbook, conn, zones: list[dict], now_ts: int,
                        window: tuple[int, int] | None = None) -> None:
    ws = wb.create_sheet("Détail par zone")

    last_col_preview = 1 + len(zones) * 3
    add_title(ws, "Détail horaire par zone", last_col_preview)
    ws["A2"] = (
        "Un relevé cumulé par heure et par zone -- pour retrouver une "
        "journée ou un pic de consommation précis. Voir « Résumé » pour la "
        "vue mensuelle."
    )
    ws["A2"].font = NOTE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col_preview)

    group_row, header_row = 3, 4
    col = 2
    zone_cols = []
    for z in zones:
        start_col = col
        for suffix in ("Consommation totale (kWh)", "dont Solaire autoconsommé (kWh)", "dont Acheté au réseau (kWh)"):
            ws.cell(row=header_row, column=col, value=suffix)
            col += 1
        ws.merge_cells(start_row=group_row, start_column=start_col, end_row=group_row, end_column=col - 1)
        gcell = ws.cell(row=group_row, column=start_col, value=z["label"])
        gcell.font, gcell.fill, gcell.alignment = GROUP_FONT, HEADER_FILL, CENTER
        zone_cols.append((z, start_col))

    ws.cell(row=header_row, column=1, value="Date et heure de la mesure")
    for c in range(1, col):
        cell = ws.cell(row=header_row, column=c)
        if cell.value is not None and cell.font != GROUP_FONT:
            cell.font, cell.fill = HEADER_FONT, SUBHEADER_FILL

    reseau_data, solaire_data = {}, {}
    all_hours = set()
    for z in zones:
        zid = z["zone"]
        for key, store in (("reseau", reseau_data), ("solaire", solaire_data)):
            src = z["sources"][key]
            if not src:
                store[zid] = {}
                continue
            rng = series_range(conn, src["series_id"])
            if not rng:
                store[zid] = {}
                continue
            q_start = max(rng[0], window[0]) if window else rng[0]
            q_end = min(rng[1], now_ts, window[1]) if window else min(rng[1], now_ts)
            data = hourly_readings(conn, src["series_id"], q_start, q_end)
            store[zid] = data
            all_hours.update(data.keys())

    r = header_row + 1
    for hour in sorted(all_hours):
        dcell = ws.cell(row=r, column=1, value=fmt_datetime(hour))
        dcell.number_format = "dd.mm.yyyy hh:mm"
        for z, start_col in zone_cols:
            zid = z["zone"]
            rv = reseau_data.get(zid, {}).get(hour)
            sv = solaire_data.get(zid, {}).get(hour)
            total = (rv or 0) + (sv or 0) if (rv is not None or sv is not None) else None
            ws.cell(row=r, column=start_col, value=round(total, 3) if total is not None else None)
            ws.cell(row=r, column=start_col + 1, value=round(sv, 3) if sv is not None else None)
            ws.cell(row=r, column=start_col + 2, value=round(rv, 3) if rv is not None else None)
        r += 1

    ws.freeze_panes = f"B{header_row + 1}"
    ws.column_dimensions["A"].width = 20
    for _z, start_col in zone_cols:
        for i in range(3):
            ws.column_dimensions[get_column_letter(start_col + i)].width = 24


# --------------------------------------------------------------------------
# Assemblage
# --------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("config", help="chemin du fichier config.yaml à utiliser (donne le db_path)")
    ap.add_argument("--miniserver", default=None,
                     help="site à exporter (nom dans config.yaml) -- requis si la config en liste plusieurs")
    ap.add_argument("--out", default=None, help="chemin du fichier .xlsx de sortie")
    ap.add_argument("--from-mois", default=None, metavar="AAAA-MM",
                     help="premier mois à inclure (défaut : tout l'historique disponible)")
    ap.add_argument("--to-mois", default=None, metavar="AAAA-MM",
                     help="dernier mois à inclure, bornes incluses (défaut : le plus récent disponible)")
    args = ap.parse_args()

    cfg = load_config(args.config)
    ms_names = [ms.name for ms in cfg.miniservers]
    if args.miniserver:
        if args.miniserver not in ms_names:
            print(f"Miniserver '{args.miniserver}' introuvable dans {args.config}. "
                  f"Sites configurés : {', '.join(ms_names)}", file=sys.stderr)
            sys.exit(1)
        ms_name = args.miniserver
    elif len(ms_names) == 1:
        ms_name = ms_names[0]
    else:
        print(f"Plusieurs sites configurés dans {args.config} ({', '.join(ms_names)}) -- "
              f"précise --miniserver NOM pour ne pas mélanger leurs zones dans le même classeur.",
              file=sys.stderr)
        sys.exit(1)

    conn = db.get_connection(cfg.db_path)
    now_ts = int(time.time())

    series = [s for s in db.list_series(conn) if s["miniserver"] == ms_name]
    zones = billing.resolve_zones(series)
    batiment_src = billing.resolve_batiment(series)

    zone_series_ids = [
        s["series_id"] for z in zones for k in ("reseau", "solaire")
        for s in [z["sources"][k]] if s
    ]
    firsts, lasts = [], []
    for sid in zone_series_ids:
        rng = series_range(conn, sid)
        if rng:
            firsts.append(rng[0])
            lasts.append(rng[1])
    if not firsts:
        print("Aucune série Réseau/Solaire par zone trouvée -- rien à exporter.", file=sys.stderr)
        sys.exit(1)

    periods = billing.periods_covering(min(firsts), max(lasts))

    window = None
    if args.from_mois or args.to_mois:
        from_key = args.from_mois or periods[0]["key"]
        to_key = args.to_mois or periods[-1]["key"]
        periods = [p for p in periods if from_key <= p["key"] <= to_key]
        if not periods:
            print(f"Aucun mois entre {from_key} et {to_key} dans les données disponibles.", file=sys.stderr)
            sys.exit(1)
        window = (periods[0]["start"], periods[-1]["end"])

    decompte = billing.compute_decompte(conn, series, periods, [], now_ts)

    wb = Workbook()
    wb.remove(wb.active)
    build_resume(wb, ms_name, decompte)
    build_legende(wb, ms_name, now_ts, periods, zones)
    build_compteur_alimentation(wb, conn, batiment_src, now_ts, window)
    build_detail_zones(wb, conn, zones, now_ts, window)

    out_path = Path(args.out) if args.out else Path("docs") / f"Mesure {ms_name}.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Classeur généré : {out_path} ({len(periods)} mois, {len(zones)} zones)")


if __name__ == "__main__":
    main()
