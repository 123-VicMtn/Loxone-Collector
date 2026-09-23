#!/usr/bin/env python3
"""
Export mensuel des compteurs d'UN SEUL appartement, destiné à être remis
au client/propriétaire tel quel (Excel + CSV).

Différence avec scripts/export_mesures_xlsx.py : celui-là exporte TOUTES les
zones d'un site, et uniquement l'électricité (Réseau/Solaire des blocs EFM
d'Arlopi). Celui-ci se concentre sur un appartement et couvre toutes les
familles de compteurs présentes sur le lot -- électricité (réseau, solaire,
total), eau (froide, chaude) et chauffage -- ce qui correspond aux postes
d'un décompte de charges complet.

Les consommations sont calculées comme un relevé de compteur classique :
valeur de l'index à la fin du mois moins valeur au début du mois (jamais une
somme de deltas horaires), en heure locale Europe/Zurich, via
billing._reading_delta -- la même primitive que /decompte, donc les chiffres
de ce classeur et ceux de la page web ne peuvent pas divergerr.

Usage :
    python3 scripts/export_appartement.py <config.yaml> --appartement APP35 \\
        [--miniserver NOM] [--from-mois AAAA-MM] [--to-mois AAAA-MM] \\
        [--out chemin.xlsx] [--csv chemin.csv]

Exemple (appartement 35 de Sequoia, janvier à mai 2026) :
    python3 scripts/export_appartement.py config.external.yaml \\
        --miniserver MS-PPE-Sequoia --appartement APP35 \\
        --from-mois 2026-01 --to-mois 2026-05

Prérequis : l'historique doit être présent en base. Le poller ne collecte
qu'à partir de son démarrage -- pour des mois passés, lancer d'abord
scripts/backfill_statistics.py sur le site concerné.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import re
import sys
import time
from pathlib import Path
from zoneinfo import ZoneInfo

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import billing
import db
import repartition
from config import load_config

TZ = ZoneInfo(billing.TIMEZONE)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FAMILY_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
NOTE_FONT = Font(italic=True, color="808080")
ALERT_FONT = Font(color="C00000")
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center")
WRAP_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)


# --------------------------------------------------------------------------
# Identification des compteurs d'un appartement
# --------------------------------------------------------------------------
# Une famille = un poste de charges. L'ordre compte : la première règle qui
# matche gagne, et les règles sur le LABEL passent avant celles sur le
# resource_type -- parce que classification.py range les compteurs de chaleur
# en "energie_consommee" (il n'a pas de type "chauffage"), donc seul le nom
# du point permet de les distinguer d'un compteur électrique.
#
# unite : imposée ici et pas lue depuis series_meta.unit, qui est vide pour
# les compteurs d'eau et de chaleur de Sequoia (Loxone ne renseigne le format
# par output que sur les blocs EFM -- voir loxone_client.
# _extract_statistic_output_units).

# Compteurs à ignorer : ce ne sont pas des mesures de consommation.
# CNrMachine est un nombre de cycles, traité par une règle dédiée -- ne
# pas l'ajouter ici.
IGNORE_RE = re.compile(r"(?i)(nombre|compteur\s+machine)")

# "Ccaleur" n'est pas une coquille de ce fichier : c'est le nom réel d'un
# point de l'appartement 34 sur MS-PPE-Sequoia ("APP 34 compteur Ccaleur
# volume [53]"). La regex tolère la variante pour ne pas classer ce compteur
# de chauffage en "contrôle" faute d'un 'h'.
CHALEUR = r"c[ch]aleur|chauffage"

RULES = [
    # (nom de la règle, regex sur le label, resource_type attendu ou None,
    #  famille, libellé affiché, unité, décimales, seuil de rupture)
    ("chauffage_energie", re.compile(rf"(?i)({CHALEUR}).*kwh"), None,
     "Chauffage", "Chauffage -- chaleur consommée", "kWh", 2, 0.5),
    ("chauffage_volume", re.compile(rf"(?i)({CHALEUR}).*volume"), None,
     "Chauffage", "Chauffage -- volume circulé", "m³", 3, 0.05),
    ("eau_chaude", None, "eau_chaude",
     "Eau", "Eau chaude sanitaire", "m³", 3, 0.05),
    ("eau_froide", None, "eau_froide",
     "Eau", "Eau froide", "m³", 3, 0.05),
    ("grid", None, "energie_reseau",
     "Électricité", "Électricité -- réseau (achetée)", "kWh", 2, 0.5),
    ("solaire", None, "energie_solaire",
     "Électricité", "Électricité -- solaire (autoconsommée)", "kWh", 2, 0.5),
    # CEnergieAppXX / CNrMachineAppXX : buanderie attribuée au lot via le
    # badge NFC (catégorie Loxone « Lessive »). Les compteurs physiques
    # M1-M6 Zähler n'ont pas d'appartement et n'entrent pas ici.
    ("buanderie", re.compile(r"(?i)cenergie"), None,
     "Buanderie", "Buanderie", "kWh", 2, 0.5),
    ("buanderie_cycles", re.compile(r"(?i)cnrmachine"), None,
     "Buanderie", "Buanderie -- cycles", "cycles", 0, 0.9),
]

# Ce qui reste en "energie_consommee" après les règles ci-dessus : un
# compteur électrique global du lot. À Arlopi ces compteurs (génération
# antérieure aux blocs EFM) mesurent un périmètre DIFFÉRENT de Grid+Solaire
# et ne servent jamais à facturer -- voir CLAUDE.md. Ils sont donc exportés
# à part, en contrôle, et jamais additionnés au reste.
CONTROLE_FAMILY = "Contrôle"
FAMILLE_ORDER = {"Électricité": 0, "Eau": 1, "Chauffage": 2, "Buanderie": 3,
                 CONTROLE_FAMILY: 4}
BUANDERIE_ROLES = {"buanderie", "buanderie_cycles"}
# Note de l'installateur Loxone : buanderie « livrée depuis le 17 juillet
# 2026 ». Avant cette date CEnergie incrémentait tous les lots (~37 kWh)
# alors que CNrMachine restait à 0 (sauf APP100). Vérifié sur les relevés :
# à partir du 17.07, 0 cycle = 0 kWh.
BUANDERIE_VALID_FROM = int(dt.datetime(2026, 7, 17, tzinfo=TZ).timestamp())
# Premier relevé Statistics du Wallbox Energizähler (index 0,000 kWh).
# Avant cette date il n'y a rien à facturer : l'index démarre à zéro.
WALLBOX_LABEL_RE = re.compile(r"(?i)wallbox\s+energiz")
WALLBOX_VALID_FROM = int(dt.datetime(2026, 3, 4, 20, 0, tzinfo=TZ).timestamp())


def _nom_court(label: str) -> str:
    """Extrait de quoi distinguer deux compteurs d'un même poste.

    "Commun - Armoire murale salle de réunion compteur chaleur kWh [2] (total)"
    -> "Armoire murale salle de réunion". On retire le préfixe de zone et la
    partie technique du nom, qui sont identiques entre les compteurs à
    distinguer et n'aident donc pas à les reconnaître.
    """
    txt = re.sub(r"\s*\((?:total)\)\s*$", "", label or "").strip()
    txt = re.sub(r"^[^-]{1,20}\s+-\s+", "", txt)
    txt = re.sub(r"\s*compteur\b.*$", "", txt, flags=re.I).strip()
    txt = re.sub(r"\s*\[\d+\]\s*$", "", txt).strip()
    return txt or (label or "").strip()


def resolve_compteurs(series: list[dict]) -> list[dict]:
    """Range les séries cumulatives ('total') d'un appartement par famille de
    charges. Retourne une liste de compteurs prêts à exporter."""
    totals = [s for s in series if s["state_name"] == "total"
              and not IGNORE_RE.search(s["label"] or "")]
    out = []
    for s in totals:
        label = s["label"] or ""
        matched = None
        for name, label_re, rtype, famille, titre, unite, dec, min_drop in RULES:
            if label_re is not None and not label_re.search(label):
                continue
            if label_re is None and s.get("resource_type") != rtype:
                continue
            if label_re is not None and rtype is not None and s.get("resource_type") != rtype:
                continue
            matched = dict(role=name, famille=famille, titre=titre, unite=unite,
                           decimales=dec, min_drop=min_drop)
            break
        if matched is None:
            if s.get("resource_type") != "energie_consommee":
                continue
            # Les sorties du bloc "Répartition Solaire" de Loxone ne sont pas
            # des compteurs ("Communs Sol" à Sequoia) : les afficher en
            # contrôle ferait croire à un relevé physique, alors que ce sont
            # les valeurs calculées dont l'incohérence est justement la
            # raison d'être de repartition.py.
            if (s.get("category") or "").strip().lower() == "répartition solaire":
                continue
            matched = dict(role="controle", famille=CONTROLE_FAMILY,
                           titre=f"{label.replace(' (total)', '')} (compteur de contrôle)",
                           unite="kWh", decimales=2, min_drop=0.5)
        # La clé doit être unique PAR SÉRIE. Utiliser le nom de la règle
        # ferait collisionner deux compteurs d'un même poste dans le dict
        # `data` (les trois compteurs de chaleur des Communs de Sequoia
        # affichaient ainsi la même valeur, et un seul des trois comptait).
        matched.update(cle=f"{matched['role']}:{s['series_id']}",
                       series_id=s["series_id"], label_source=label)
        out.append(matched)

    # Deux compteurs d'un même poste (trois compteurs de chaleur dans les
    # communs) portent le même libellé : on les distingue par leur propre nom,
    # sinon le classeur montre plusieurs lignes indiscernables.
    par_titre = {}
    for c in out:
        par_titre.setdefault(c["titre"], []).append(c)
    for titre, group in par_titre.items():
        if len(group) < 2:
            continue
        for c in group:
            court = _nom_court(c["label_source"])
            if court and court.lower() not in titre.lower():
                c["titre"] = f"{titre} -- {court}"

    out.sort(key=lambda c: (FAMILLE_ORDER.get(c["famille"], 9), c["titre"]))
    return out


# --------------------------------------------------------------------------
# Calcul
# --------------------------------------------------------------------------

def compute(conn, compteurs: list[dict], periods: list[dict], now_ts: int) -> dict:
    """Pour chaque compteur et chaque mois : index de début, index de fin,
    consommation, et alertes éventuelles."""
    data = {}
    periode_couvre_buanderie = periods[-1]["end"] > BUANDERIE_VALID_FROM
    for c in compteurs:
        per_month = {}
        for p in periods:
            start, end = p["start"], p["end"]
            notes = []
            if c.get("role") in BUANDERIE_ROLES and periode_couvre_buanderie:
                if end <= BUANDERIE_VALID_FROM:
                    per_month[p["key"]] = {
                        "kwh": 0.0, "releve_debut": None, "releve_fin": None,
                        "alertes": ["avant mise en service buanderie (17.07.2026)"],
                        "en_cours": start <= now_ts < end,
                    }
                    continue
                if start < BUANDERIE_VALID_FROM:
                    start = BUANDERIE_VALID_FROM
                    notes.append("buanderie décomptée depuis le 17.07.2026 "
                                 "(mise en service Loxone)")
            if WALLBOX_LABEL_RE.search(c.get("label_source") or ""):
                if end <= WALLBOX_VALID_FROM:
                    per_month[p["key"]] = {
                        "kwh": 0.0, "releve_debut": None, "releve_fin": None,
                        "alertes": ["avant historique Wallbox (04.03.2026)"],
                        "en_cours": start <= now_ts < end,
                    }
                    continue
                if start < WALLBOX_VALID_FROM:
                    start = WALLBOX_VALID_FROM
                    notes.append("historique Wallbox à partir du 04.03.2026")
            d = billing._reading_delta(conn, c["series_id"], start, end,
                                       now_ts, min_drop=c["min_drop"])
            d["en_cours"] = p["start"] <= now_ts < p["end"]
            if notes:
                d["alertes"] = notes + d.get("alertes", [])
            per_month[p["key"]] = d
        data[c["cle"]] = per_month

    # Consommation électrique totale = réseau + solaire autoconsommé, la
    # sémantique établie empiriquement sur les blocs EFM (voir CLAUDE.md).
    # Calculée seulement si les DEUX compteurs existent : sur un site sans
    # split par zone (PPE-Horizon), la ligne n'a pas de sens et est omise.
    roles = {c["role"]: c["cle"] for c in compteurs}
    if "grid" in roles and "solaire" in roles:
        per_month = {}
        for p in periods:
            g = data[roles["grid"]][p["key"]]
            s = data[roles["solaire"]][p["key"]]
            total = None
            if g["kwh"] is not None and s["kwh"] is not None:
                total = g["kwh"] + s["kwh"]
            per_month[p["key"]] = {
                "kwh": total,
                "releve_debut": None,
                "releve_fin": None,
                "alertes": [],
                "en_cours": g["en_cours"],
                "taux_solaire": (100.0 * s["kwh"] / total) if (total and total > 0
                                                               and s["kwh"] is not None) else None,
            }
        data["electricite_totale"] = per_month
    return data


def appliquer_repartition_calculee(conn, compteurs: list[dict], data: dict,
                                   periods: list[dict], now_ts: int,
                                   ev: dict, conso_src: dict) -> list[str]:
    """Remplace les lignes réseau/solaire venues du Miniserver par une
    répartition reconstruite (voir repartition.py), quand les séries du
    Miniserver ont été jugées incohérentes.

    Le TOTAL facturé reste le relevé du compteur du lot -- index de fin moins
    index de début, via billing._reading_delta, comme tous les autres postes
    du classeur. Seule la PROPORTION réseau/solaire vient de l'agrégation
    horaire. Faire autrement (sommer les deltas horaires) donnerait un total
    légèrement différent de celui de l'onglet « Relevés de compteur », donc
    un classeur qui ne se recalcule pas à la main.
    """
    notes = []
    frac = ev["fraction"]["fraction"]

    # Les lignes réseau/solaire du Miniserver sortent du classeur : les
    # laisser à côté des lignes recalculées inviterait à les additionner.
    for c in [x for x in compteurs if x.get("role") in ("grid", "solaire")]:
        compteurs.remove(c)
        data.pop(c["cle"], None)
    data.pop("electricite_totale", None)

    # Le compteur du lot passe de "contrôle" à poste principal : c'est lui
    # qui mesure la consommation électrique facturable.
    total_cle = None
    for c in compteurs:
        if c.get("series_id") == conso_src["series_id"]:
            c.update(famille="Électricité", decimales=2, unite="kWh",
                     role="electricite_totale",
                     titre="Électricité -- consommation totale (compteur du lot)",
                     highlight=True, avec_taux=True)
            total_cle = c["cle"]
    if total_cle is None:
        # Le compteur du lot n'était pas dans les séries de l'appartement
        # (classification à revoir) : sans lui, rien à répartir.
        notes.append("compteur de consommation du lot introuvable : "
                     "répartition réseau/solaire impossible.")
        return notes

    reseau_mois, solaire_mois = {}, {}
    non_repartis = 0.0
    for p in periods:
        mesure = data[total_cle][p["key"]]
        total = mesure["kwh"]
        split = repartition.split_zone(conn, conso_src["series_id"], frac,
                                       p["start"], p["end"])
        # Part solaire déduite des seules heures dont on connaît le mix.
        base = split["total"] - split["non_reparti"]
        part = (split["solaire"] / base) if base > 0 else None
        non_repartis += split["non_reparti"]

        sol = res = None
        if total is not None and part is not None:
            # Arrondis à la précision AFFICHÉE, et part réseau déduite de la
            # part solaire arrondie : sinon les deux lignes du classeur
            # peuvent afficher 11,76 + 2,34 = 14,10 en face d'un total de
            # 14,09, ce qui contredit l'égalité exacte promise dans le
            # « Lisez-moi ».
            sol = round(total * part, 2)
            total = round(total, 2)
            res = total - sol
            # Le total mensuel est figé sur la même précision : sans ça, la
            # colonne « Total période » sommerait des valeurs non arrondies
            # (66,85) face à des parts arrondies (33,60 + 33,24 = 66,84).
            mesure["kwh"] = total
        elif total is not None and base <= 0 and total <= 0:
            sol, res = 0.0, 0.0

        commun = {"releve_debut": None, "releve_fin": None, "alertes": [],
                  "en_cours": mesure["en_cours"]}
        reseau_mois[p["key"]] = dict(commun, kwh=res)
        solaire_mois[p["key"]] = dict(commun, kwh=sol)
        mesure["taux_solaire"] = (100.0 * part) if part is not None else None

    data["grid"] = reseau_mois
    data["solaire"] = solaire_mois
    compteurs.append(dict(cle="grid", role="grid", famille="Électricité",
                          titre="Électricité -- réseau (achetée)", unite="kWh",
                          decimales=2, min_drop=0.5, calcule=True))
    compteurs.append(dict(cle="solaire", role="solaire", famille="Électricité",
                          titre="Électricité -- solaire (autoconsommée)", unite="kWh",
                          decimales=2, min_drop=0.5, calcule=True))

    compteurs.sort(key=lambda c: (FAMILLE_ORDER.get(c["famille"], 9), c["titre"]))

    if non_repartis > 0.5:
        notes.append(f"{non_repartis:.1f} kWh consommés pendant des heures sans "
                     f"mesure de bâtiment : leur part solaire est estimée d'après "
                     f"le reste de la période.")
    return notes


def total_periode(per_month: dict, periods: list[dict]) -> float | None:
    """Somme sur la plage exportée -- None dès qu'un mois manque, pour ne
    jamais présenter un total silencieusement incomplet."""
    acc = 0.0
    for p in periods:
        v = per_month[p["key"]]["kwh"]
        if v is None:
            return None
        acc += v
    return acc


# --------------------------------------------------------------------------
# Rendu Excel
# --------------------------------------------------------------------------

def _milliers(v: float) -> str:
    """21681 -> '21 681' : espace insécable fine, comme sur une facture suisse."""
    return f"{v:,.0f}".replace(",", "\u202f")


def add_title(ws, text: str, last_col: int, row: int = 1) -> None:
    ws.cell(row=row, column=1, value=text).font = TITLE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(last_col, 1))


def build_lisezmoi(wb: Workbook, site: str, apt_label: str, compteurs: list[dict],
                   periods: list[dict], now_ts: int, alertes: list[str],
                   data: dict | None = None, ev: dict | None = None) -> None:
    ws = wb.create_sheet("Lisez-moi")
    add_title(ws, f"Relevés de consommation -- {apt_label}", 2)

    date_str = dt.datetime.fromtimestamp(now_ts, TZ).strftime("%d.%m.%Y à %H:%M")
    periode = f"{periods[0]['label']} -- {periods[-1]['label']}" if periods else "-"

    rows = [
        ("", ""),
        ("Appartement", apt_label),
        ("Bâtiment (installation)", site),
        ("Période couverte", periode),
        ("Généré le", f"{date_str} (heure de Suisse)"),
        ("", ""),
        ("Comment ces chiffres sont obtenus", ""),
        ("Méthode",
         "Chaque consommation mensuelle est la différence entre l'index du "
         "compteur à la fin du mois et son index au début du mois -- "
         "exactement comme un relevé de compteur physique à deux dates. "
         "Les index utilisés sont visibles dans l'onglet « Relevés de compteur »."),
        ("Fuseau horaire",
         "Les mois commencent et finissent à minuit, heure locale suisse "
         "(Europe/Zurich), changements d'heure inclus."),
        ("", ""),
        ("Les postes mesurés", ""),
    ]
    familles = {}
    for c in compteurs:
        familles.setdefault(c["famille"], []).append(c)
    descriptions = {
        "Électricité -- réseau (achetée)":
            "Électricité achetée au fournisseur et consommée par l'appartement.",
        "Électricité -- solaire (autoconsommée)":
            "Électricité produite par les panneaux solaires du bâtiment et "
            "consommée directement par l'appartement.",
        "Eau froide": "Volume d'eau froide consommé.",
        "Eau chaude sanitaire": "Volume d'eau chaude sanitaire consommé.",
        "Chauffage -- chaleur consommée":
            "Énergie thermique consommée pour le chauffage.",
        "Chauffage -- volume circulé":
            "Volume d'eau de chauffage passé dans le compteur (donnée technique, "
            "généralement pas facturée telle quelle).",
    }
    calculee = bool(ev and ev.get("source_recommandee") == "calculee")
    if calculee:
        descriptions.update({
            "Électricité -- consommation totale (compteur du lot)":
                "Consommation électrique de l'appartement, relevée sur son "
                "compteur. C'est le total facturable, et c'est une mesure "
                "directe.",
            "Électricité -- réseau (achetée)":
                "Part de cette consommation couverte par l'électricité "
                "achetée au réseau. Répartition calculée (voir ci-dessous).",
            "Électricité -- solaire (autoconsommée)":
                "Part de cette consommation couverte par la production "
                "solaire du bâtiment. Répartition calculée (voir ci-dessous).",
        })
    for famille in ("Électricité", "Eau", "Chauffage", CONTROLE_FAMILY):
        for c in familles.get(famille, []):
            desc = descriptions.get(c["titre"])
            if c["famille"] == CONTROLE_FAMILY:
                desc = ("Compteur électrique global du lot, conservé à titre de "
                        "contrôle. Son périmètre de mesure diffère de la somme "
                        "réseau + solaire : il n'est pas utilisé pour la facturation.")
            rows.append((f"{c['titre']} ({c['unite']})", desc or ""))
    if not calculee:
        rows += [
            ("", ""),
            ("Consommation électrique totale",
             "Somme du réseau et du solaire autoconsommé : tout ce que "
             "l'appartement a réellement consommé en électricité."),
        ]
    else:
        f = ev["fraction"]
        rows += [
            ("", ""),
            ("Comment la part réseau et la part solaire sont calculées", ""),
            ("Principe",
             "Le compteur de l'appartement mesure sa consommation totale, "
             "mais pas son origine. La part solaire est donc calculée heure "
             "par heure : pour chaque heure, on détermine quelle proportion "
             "de l'électricité consommée dans le bâtiment venait des "
             "panneaux solaires, et on applique cette proportion à la "
             "consommation de l'appartement sur cette même heure. C'est la "
             "méthode utilisée par les prestataires de regroupement de "
             "consommation propre (RCP)."),
            ("Formule, pour chaque heure",
             "part solaire du bâtiment = (production solaire - électricité "
             "réinjectée au réseau) / (électricité achetée au réseau + "
             "production solaire - électricité réinjectée). "
             "La part réinjectée est retirée : revendue au réseau, elle n'a "
             "pas été consommée dans l'immeuble."),
            ("Garantie",
             "Part réseau + part solaire redonne exactement la consommation "
             "relevée au compteur de l'appartement. La répartition ne fait "
             "que scinder ce total en deux origines ; elle n'ajoute et ne "
             "retire aucun kWh."),
            ("Production du bâtiment sur la période",
             f"{_milliers(f['production'])} kWh produits, dont "
             f"{_milliers(f['autoconsomme'])} kWh consommés dans l'immeuble et "
             f"{_milliers(f['injection'])} kWh réinjectés au réseau. "
             f"Électricité achetée au réseau : "
             f"{_milliers(f['import_reseau'])} kWh."),
            ("Résolution du calcul",
             "Horaire, c'est la finesse de l'historique enregistré par "
             "l'installation. Un décompte au quart d'heure pourrait donner "
             "un écart de l'ordre du pourcent sur la répartition, sans "
             "changer le total consommé."),
        ]
    rows += [
        ("", ""),
        ("Onglets du classeur", ""),
        ("« Résumé mensuel »", "Un chiffre par mois et par compteur -- la vue à retenir."),
        ("« Relevés de compteur »",
         "Les index de début et de fin de chaque mois, pour pouvoir refaire "
         "le calcul à la main."),
    ]

    # Total de la période, poste par poste. Sur un lot peu utilisé, les
    # volumes d'eau en m^3 s'affichent "0,004" dans le Résumé et se lisent
    # comme une donnée manquante : on répète donc la valeur en litres, qui
    # montre sans ambiguïté que le compteur a bien mesuré quelque chose.
    if data:
        rows += [("", ""), ("Total sur la période", "")]
        for c in compteurs:
            if c["famille"] == CONTROLE_FAMILY:
                continue
            tot = total_periode(data[c["cle"]], periods)
            if tot is None:
                rows.append((c["titre"], "Non calculable (données incomplètes)."))
                continue
            txt = f"{tot:.{c['decimales']}f} {c['unite']}".replace(".", ",")
            # Le rappel en litres ne sert que pour les très petits volumes,
            # ceux qui s'affichent "0,004" et se lisent comme une absence de
            # mesure. Au-delà, il rallonge sans rien clarifier.
            if c["unite"] == "m³" and tot < 10:
                txt += f"  (soit {tot * 1000:.0f} litres)"
            rows.append((c["titre"], txt))
    if alertes:
        rows += [("", ""), ("Points d'attention", "")]
        rows += [("", a) for a in alertes]

    r = 3
    for label, text in rows:
        if label and not text:
            ws.cell(row=r, column=1, value=label).font = BOLD
        else:
            ws.cell(row=r, column=1, value=label).font = BOLD
            cell = ws.cell(row=r, column=2, value=text)
            cell.alignment = WRAP_LEFT
        r += 1

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 95


def build_resume(wb: Workbook, apt_label: str, compteurs: list[dict], data: dict,
                 periods: list[dict], now_ts: int) -> None:
    ws = wb.create_sheet("Résumé mensuel")
    last_col = 3 + len(periods)
    add_title(ws, f"Consommations mensuelles -- {apt_label}", last_col)
    ws["A2"] = (
        "Consommation de chaque mois = index de fin de mois - index de début de mois. "
        "Une cellule vide signifie qu'aucune donnée fiable n'est disponible pour ce mois."
    )
    ws["A2"].font = NOTE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)

    header_row = 4
    ws.cell(row=header_row, column=1, value="Poste")
    ws.cell(row=header_row, column=2, value="Unité")
    for i, p in enumerate(periods):
        ws.cell(row=header_row, column=3 + i, value=p["label"])
    total_col = 3 + len(periods)
    ws.cell(row=header_row, column=total_col, value="Total période")
    for c in range(1, total_col + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, CENTER

    familles = {}
    for c in compteurs:
        familles.setdefault(c["famille"], []).append(c)

    r = header_row + 1
    for famille in ("Électricité", "Eau", "Chauffage", CONTROLE_FAMILY):
        items = familles.get(famille)
        if not items:
            continue
        fcell = ws.cell(row=r, column=1, value=famille)
        fcell.font = BOLD
        for c in range(1, total_col + 1):
            ws.cell(row=r, column=c).fill = FAMILY_FILL
        r += 1

        for c in items:
            r = _write_row(ws, r, c["titre"], c["unite"], c["decimales"],
                           data[c["cle"]], periods, total_col,
                           highlight=c.get("highlight", False))
            # En répartition recalculée, le total est un compteur réel (celui
            # du lot) et porte donc lui-même la ligne de part solaire.
            if c.get("avec_taux"):
                r = _write_taux(ws, r, data[c["cle"]], periods, total_col)

        # La ligne de synthèse électrique se place juste après réseau/solaire.
        if famille == "Électricité" and "electricite_totale" in data:
            r = _write_row(ws, r, "Consommation électrique totale", "kWh", 2,
                           data["electricite_totale"], periods, total_col,
                           highlight=True)
            r = _write_taux(ws, r, data["electricite_totale"], periods, total_col)
        r += 1

    ws.freeze_panes = f"C{header_row + 1}"
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 8
    for i in range(len(periods) + 1):
        ws.column_dimensions[get_column_letter(3 + i)].width = 13
    return


def _write_row(ws, r: int, titre: str, unite: str, decimales: int, per_month: dict,
               periods: list[dict], total_col: int, highlight: bool = False) -> int:
    ws.cell(row=r, column=1, value=titre)
    ws.cell(row=r, column=2, value=unite).alignment = CENTER
    fmt = "0." + "0" * decimales if decimales else "0"
    for i, p in enumerate(periods):
        e = per_month[p["key"]]
        val = e["kwh"]
        cell = ws.cell(row=r, column=3 + i,
                       value=round(val, decimales) if val is not None else None)
        cell.number_format = fmt
        if e.get("en_cours"):
            cell.font = NOTE_FONT
        elif val is None:
            cell.font = ALERT_FONT
    tot = total_periode(per_month, periods)
    tcell = ws.cell(row=r, column=total_col,
                    value=round(tot, decimales) if tot is not None else None)
    tcell.number_format = fmt
    tcell.font = BOLD
    if highlight:
        for c in range(1, total_col + 1):
            ws.cell(row=r, column=c).fill = TOTAL_FILL
        ws.cell(row=r, column=1).font = BOLD
    return r + 1


def _write_taux(ws, r: int, per_month: dict, periods: list[dict], total_col: int) -> int:
    ws.cell(row=r, column=1, value="  dont part solaire")
    ws.cell(row=r, column=2, value="%").alignment = CENTER
    for i, p in enumerate(periods):
        val = per_month[p["key"]].get("taux_solaire")
        cell = ws.cell(row=r, column=3 + i, value=round(val, 1) if val is not None else None)
        cell.number_format = "0.0"
    return r + 1


def build_releves(wb: Workbook, apt_label: str, compteurs: list[dict], data: dict,
                  periods: list[dict]) -> None:
    ws = wb.create_sheet("Relevés de compteur")
    last_col = 6
    add_title(ws, f"Index de compteur début / fin de mois -- {apt_label}", last_col)
    ws["A2"] = (
        "Les index bruts ayant servi au calcul : consommation = index de fin - index de début. "
        "La date indiquée est celle du relevé effectivement disponible le plus proche de la borne du mois."
    )
    ws["A2"].font = NOTE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)

    header_row = 4
    headers = ["Mois", "Poste", "Unité", "Index début", "Index fin", "Consommation"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=i, value=h)
        cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, CENTER

    r = header_row + 1
    for p in periods:
        for c in compteurs:
            e = data[c["cle"]][p["key"]]
            fmt = "0." + "0" * c["decimales"] if c["decimales"] else "0"
            ws.cell(row=r, column=1, value=p["label"])
            ws.cell(row=r, column=2, value=c["titre"])
            ws.cell(row=r, column=3, value=c["unite"]).alignment = CENTER
            for col, key in ((4, "releve_debut"), (5, "releve_fin"), (6, "kwh")):
                val = e[key]
                cell = ws.cell(row=r, column=col,
                               value=round(val, c["decimales"]) if val is not None else None)
                cell.number_format = fmt
            if e["alertes"]:
                ws.cell(row=r, column=7, value=" ; ".join(e["alertes"])).font = ALERT_FONT
            r += 1

    ws.freeze_panes = f"A{header_row + 1}"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 8
    for col in ("D", "E", "F"):
        ws.column_dimensions[col].width = 15
    ws.column_dimensions["G"].width = 60


# --------------------------------------------------------------------------
# Rendu CSV
# --------------------------------------------------------------------------

def write_csv(path: Path, apt_label: str, site: str, compteurs: list[dict],
              data: dict, periods: list[dict]) -> None:
    """CSV plat, une ligne par mois et par poste -- pour retraitement dans un
    tableur ou un logiciel de gérance. Séparateur ';' et virgule décimale :
    Excel en locale FR/CH ouvre le fichier directement, sans assistant
    d'import."""
    rows = []
    for p in periods:
        for c in compteurs:
            e = data[c["cle"]][p["key"]]
            remarques = list(e["alertes"])
            # Une ligne sans index de compteur doit dire d'où elle vient,
            # sinon elle est indiscernable d'une mesure directe.
            if c.get("calcule"):
                remarques.insert(0, "répartition calculée (pas un relevé de compteur)")
            rows.append({
                "appartement": apt_label,
                "batiment": site,
                "mois": p["key"],
                "mois_libelle": p["label"],
                "poste": c["titre"],
                "unite": c["unite"],
                "index_debut": e["releve_debut"],
                "index_fin": e["releve_fin"],
                "consommation": e["kwh"],
                "remarques": " ; ".join(remarques),
            })
        if "electricite_totale" in data:
            e = data["electricite_totale"][p["key"]]
            rows.append({
                "appartement": apt_label,
                "batiment": site,
                "mois": p["key"],
                "mois_libelle": p["label"],
                "poste": "Consommation électrique totale (réseau + solaire)",
                "unite": "kWh",
                "index_debut": None,
                "index_fin": None,
                "consommation": e["kwh"],
                "remarques": "",
            })

    def fr(v, dec):
        if v is None:
            return ""
        return f"{round(v, dec):.{dec}f}".replace(".", ",")

    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Appartement", "Bâtiment", "Mois", "Mois (libellé)", "Poste",
                    "Unité", "Index début", "Index fin", "Consommation", "Remarques"])
        for row in rows:
            dec = 3 if row["unite"] == "m³" else 2
            w.writerow([row["appartement"], row["batiment"], row["mois"],
                        row["mois_libelle"], row["poste"], row["unite"],
                        fr(row["index_debut"], dec), fr(row["index_fin"], dec),
                        fr(row["consommation"], dec), row["remarques"]])


# --------------------------------------------------------------------------
# Assemblage
# --------------------------------------------------------------------------

def parse_mois(value: str) -> tuple[int, int]:
    try:
        y, m = value.split("-")
        return int(y), int(m)
    except Exception:
        raise SystemExit(f"Format de mois invalide : '{value}' (attendu AAAA-MM)")


def construire_periodes(y1: int, m1: int, y2: int, m2: int) -> list[dict]:
    periods = []
    y, m = y1, m1
    while (y, m) <= (y2, m2):
        start, end = billing.period_bounds(y, m)
        periods.append({"key": billing.period_key(y, m), "label": billing.period_label(y, m),
                        "start": start, "end": end})
        m += 1
        if m > 12:
            y, m = y + 1, 1
    return periods


def traiter_appartement(conn, ms_name: str, series_site: list[dict], zones: list[dict],
                        ev: dict, apt: str, periods: list[dict], now_ts: int,
                        out: str | None = None, csv_path: str | None = None,
                        outdir: Path | None = None, verbose: bool = True,
                        write_files: bool = True) -> dict | None:
    """Génère le classeur et le CSV d'un appartement.

    Retourne un résumé (postes -> total sur la période) pour le
    récapitulatif consolidé, ou None si la zone n'a aucun compteur
    exploitable -- cas des zones fantômes créées par une convention de
    nommage (APP100 sur Sequoia, issu des points de buanderie
    "CEnergieApp100"), qu'il ne faut pas transformer en classeur vide.
    """
    if apt == "APP100":
        if verbose:
            print(f"  (ignoré : {apt} -- zone fantôme des points de buanderie, "
                  f"pas un logement)")
        return None
    series = [s for s in series_site if (s["apartment"] or "").upper() == apt]
    compteurs = resolve_compteurs(series)
    if not compteurs:
        if verbose:
            print(f"  (ignoré : {apt} -- aucun compteur cumulatif)")
        return None

    data = compute(conn, compteurs, periods, now_ts)

    notes_repartition = []
    if ev["source_recommandee"] == "calculee":
        conso_src = next((z["sources"]["controle"] for z in zones
                          if z["zone"].upper() == apt and z["sources"].get("controle")), None)
        if conso_src:
            notes_repartition = appliquer_repartition_calculee(
                conn, compteurs, data, periods, now_ts, ev, conso_src)
        else:
            notes_repartition = ["compteur de consommation du lot introuvable : "
                                 "les séries réseau/solaire du Miniserver sont "
                                 "conservées malgré leur incohérence."]

    # Le tri des compteurs de contrôle vient APRÈS la répartition : le
    # compteur du lot y est promu en poste facturable, et l'écarter avant
    # (il est muet sur le premier mois de certaines périodes) rendrait la
    # répartition impossible sur toute la plage.
    #
    # Un compteur de CONTRÔLE dont on ne peut pas calculer le total de la
    # période n'apporte rien : soit il est muet, soit il n'existe que depuis
    # le démarrage du collecteur. Les postes FACTURABLES, eux, restent
    # toujours affichés même sans donnée -- une absence de mesure sur un
    # poste facturé doit se voir plutôt que disparaître.
    #
    # Exception : la buanderie (CEnergie) n'a aucun historique avant juin
    # 2026. Sur une période antérieure elle n'a rien à afficher, donc on
    # l'ôte plutôt que de créer une colonne entièrement vide.
    muets = [c for c in compteurs
             if (c["famille"] == CONTROLE_FAMILY
                 or c.get("role") in BUANDERIE_ROLES)
             and total_periode(data[c["cle"]], periods) is None]
    for c in muets:
        compteurs.remove(c)
        data.pop(c["cle"], None)

    # Plus aucun compteur exploitable : zone fantôme née d'une convention de
    # nommage (APP100 sur Sequoia vient des points de buanderie
    # "CEnergieApp100"), pas un logement. Pas de classeur pour elle.
    if not compteurs:
        if verbose:
            print(f"  (ignoré : {apt} -- aucun compteur exploitable sur la période)")
        return None

    apt_label = billing._zone_label(apt)
    alertes = []
    for c in compteurs:
        manquants = [p["label"] for p in periods if data[c["cle"]][p["key"]]["kwh"] is None]
        if manquants:
            alertes.append(f"{c['titre']} : aucune donnée exploitable pour "
                           f"{', '.join(manquants)}.")

    alertes += notes_repartition

    if verbose:
        print(f"\n{apt_label} ({apt}) -- {len(compteurs)} compteur(s)")
        for c in compteurs:
            vals = [data[c["cle"]][p["key"]]["kwh"] for p in periods]
            ok = sum(1 for v in vals if v is not None)
            tot = total_periode(data[c["cle"]], periods)
            tot_s = f"{tot:.{c['decimales']}f} {c['unite']}" if tot is not None else "incomplet"
            print(f"  - {c['titre']:44} [{c['famille']:12}] "
                  f"{ok}/{len(periods)} mois, total {tot_s}")
        for c in muets:
            print(f"  (écarté : {c['titre']} -- aucune donnée sur la période)")
        for a in alertes:
            print(f"  ! {a}")

    suffix = f"{apt_label} {periods[0]['key']} a {periods[-1]['key']}"
    base = outdir if outdir is not None else Path("docs")
    out_path = Path(out) if out else base / f"Consommations {suffix}.xlsx"
    cpath = Path(csv_path) if csv_path else out_path.with_suffix(".csv")
    if write_files:
        wb = Workbook()
        wb.remove(wb.active)
        build_resume(wb, apt_label, compteurs, data, periods, now_ts)
        build_releves(wb, apt_label, compteurs, data, periods)
        build_lisezmoi(wb, ms_name, apt_label, compteurs, periods, now_ts,
                       alertes, data, ev)
        wb.move_sheet("Lisez-moi", offset=-2)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out_path)
        write_csv(cpath, apt_label, ms_name, compteurs, data, periods)
        if verbose:
            print(f"  -> {out_path.name}  +  {cpath.name}")

    def poste(titre, famille, unite, dec, cle, role, calcule=False):
        return {"titre": titre, "famille": famille, "unite": unite, "decimales": dec,
                "role": role, "calcule": calcule,
                "total": total_periode(data[cle], periods),
                "mois": {p["key"]: data[cle][p["key"]]["kwh"] for p in periods}}

    postes = [poste(c["titre"], c["famille"], c["unite"], c["decimales"], c["cle"],
                    c.get("role", "controle"), bool(c.get("calcule")))
              for c in compteurs]
    # La consommation électrique totale est une ligne calculée du classeur
    # (réseau + solaire) et non un compteur : elle doit quand même figurer au
    # récapitulatif, c'est la colonne principale d'un décompte.
    if "electricite_totale" in data:
        postes.insert(0, poste("Électricité -- consommation totale", "Électricité",
                               "kWh", 2, "electricite_totale", "electricite_totale",
                               calcule=True))

    return {"apt": apt, "label": apt_label, "xlsx": out_path, "csv": cpath,
            "alertes": alertes, "postes": postes}


# Postes du décompte. Un poste regroupe TOUS les compteurs d'un même rôle
# dans une zone : les communs de Sequoia ont trois compteurs de chaleur, qui
# n'ont pas à apparaître séparément sur un document de facturation.
#   (rôle, libellé, unité, décimales, facturable)
# "facturable" = le poste reçoit un prix unitaire et entre dans le montant.
# La consommation électrique totale n'est pas facturable en tant que telle :
# elle est déjà facturée par ses deux composantes réseau et solaire, et ne
# figure que comme contrôle (elle doit égaler leur somme).
CATEGORIES = [
    ("grid", "Électricité réseau", "kWh", 2, True),
    ("solaire", "Électricité solaire", "kWh", 2, True),
    ("electricite_totale", "Électricité totale", "kWh", 2, False),
    ("eau_chaude", "Eau chaude", "m³", 3, True),
    ("eau_froide", "Eau froide", "m³", 3, True),
    ("chauffage_energie", "Chauffage", "kWh", 2, True),
    ("chauffage_volume", "Chauffage (volume)", "m³", 3, False),
    ("buanderie", "Buanderie", "kWh", 2, True),
    ("buanderie_cycles", "Buanderie -- cycles", "cycles", 0, True),
]

EDIT_FILL = PatternFill("solid", fgColor="FFF9C4")
EDIT_BORDER = Border(*[Side(style="thin", color="BFBFBF")] * 4)
MONEY_FMT = '#,##0.00'


def build_recap(resumes: list[dict], site: str, periods: list[dict], ev: dict,
                now_ts: int, out_path: Path, tva: float = 8.1) -> None:
    """Décompte consolidé : une ligne par appartement, une colonne par poste,
    et des prix unitaires modifiables qui calculent les montants.

    Volontairement sans détail horaire : les quantités sont des relevés de
    compteur sur la période, c'est tout ce qu'il faut pour facturer. Le
    détail mensuel reste dans un second onglet pour vérifier une ligne.
    """
    wb = Workbook()
    wb.remove(wb.active)

    # On ne garde que les catégories réellement présentes dans les données.
    def qte(res, role):
        vals = [p["total"] for p in res["postes"]
                if p["role"] == role and p["total"] is not None]
        return sum(vals) if vals else None

    cats = [c for c in CATEGORIES
            if any(qte(res, c[0]) is not None for res in resumes)]

    ws = wb.create_sheet("Décompte")
    col_qte0 = 2                       # première colonne de quantité
    col_ht = col_qte0 + len(cats)      # montant HT
    ncol = col_ht + 2                  # + TVA + TTC
    add_title(ws, f"Décompte de charges -- {site}", ncol)
    ws["A2"] = (f"Période : {periods[0]['label']} -- {periods[-1]['label']} "
                f"({len(periods)} mois). Quantités relevées aux compteurs. "
                f"Saisis les prix unitaires dans les cases jaunes : les montants "
                f"se calculent tout seuls.")
    ws["A2"].font = NOTE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)

    # --- Bloc tarifs, cases modifiables référencées par les formules -------
    r = 4
    ws.cell(row=r, column=1, value="Prix unitaires à compléter").font = Font(
        bold=True, size=12, color="1F4E78")
    r += 1
    prix_ref = {}
    for role, lib, unite, dec, facturable in cats:
        if not facturable:
            continue
        ws.cell(row=r, column=1, value=lib).font = BOLD
        cell = ws.cell(row=r, column=2, value=0.0)
        cell.fill, cell.border, cell.number_format = EDIT_FILL, EDIT_BORDER, MONEY_FMT
        ws.cell(row=r, column=3, value=f"CHF / {unite}").font = NOTE_FONT
        prix_ref[role] = f"$B${r}"
        r += 1
    ws.cell(row=r, column=1, value="TVA").font = BOLD
    cell = ws.cell(row=r, column=2, value=tva)
    cell.fill, cell.border, cell.number_format = EDIT_FILL, EDIT_BORDER, '0.0"%"'
    tva_ref = f"$B${r}"
    r += 2

    # --- Tableau par appartement -------------------------------------------
    hr = r
    ws.cell(row=hr, column=1, value="Appartement")
    for i, (role, lib, unite, dec, fact) in enumerate(cats):
        ws.cell(row=hr, column=col_qte0 + i, value=f"{lib}\n({unite})")
    ws.cell(row=hr, column=col_ht, value="Montant HT\n(CHF)")
    ws.cell(row=hr, column=col_ht + 1, value="TVA\n(CHF)")
    ws.cell(row=hr, column=col_ht + 2, value="Montant TTC\n(CHF)")
    for c in range(1, ncol + 1):
        cell = ws.cell(row=hr, column=c)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[hr].height = 34

    r = hr + 1
    first_data = r
    for res in resumes:
        ws.cell(row=r, column=1, value=res["label"])
        termes = []
        for i, (role, lib, unite, dec, fact) in enumerate(cats):
            v = qte(res, role)
            cell = ws.cell(row=r, column=col_qte0 + i,
                           value=round(v, dec) if v is not None else None)
            cell.number_format = "0" if dec == 0 else "0." + "0" * dec
            if fact and v is not None:
                termes.append(f"{get_column_letter(col_qte0 + i)}{r}*{prix_ref[role]}")
        ht = f"={'+'.join(termes)}" if termes else 0
        ws.cell(row=r, column=col_ht, value=ht).number_format = MONEY_FMT
        ws.cell(row=r, column=col_ht + 1,
                value=f"={get_column_letter(col_ht)}{r}*{tva_ref}/100"
                ).number_format = MONEY_FMT
        ws.cell(row=r, column=col_ht + 2,
                value=f"={get_column_letter(col_ht)}{r}+{get_column_letter(col_ht + 1)}{r}"
                ).number_format = MONEY_FMT
        r += 1

    ws.cell(row=r, column=1, value="TOTAL").font = BOLD
    for c in range(col_qte0, ncol + 1):
        L = get_column_letter(c)
        cell = ws.cell(row=r, column=c, value=f"=SUM({L}{first_data}:{L}{r - 1})")
        dec = cats[c - col_qte0][3] if c < col_ht else 2
        cell.number_format = (MONEY_FMT if c >= col_ht
                              else ("0" if dec == 0 else "0." + "0" * dec))
        cell.font = BOLD
    for c in range(1, ncol + 1):
        ws.cell(row=r, column=c).fill = TOTAL_FILL

    ws.freeze_panes = f"B{hr + 1}"
    ws.column_dimensions["A"].width = 22
    for c in range(col_qte0, ncol + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14

    # Détail mensuel, une ligne par appartement / poste / mois.
    wd = wb.create_sheet("Détail mensuel")
    cols = ["Appartement", "Poste", "Unité"] + [p["label"] for p in periods] + ["Total"]
    add_title(wd, f"Détail mensuel -- {site}", len(cols))
    for i, h in enumerate(cols, start=1):
        cell = wd.cell(row=3, column=i, value=h)
        cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, CENTER
    rr = 4
    for res in resumes:
        for p in res["postes"]:
            fmt = "0." + "0" * p["decimales"] if p["decimales"] else "0"
            wd.cell(row=rr, column=1, value=res["label"])
            wd.cell(row=rr, column=2, value=p["titre"])
            wd.cell(row=rr, column=3, value=p["unite"]).alignment = CENTER
            for i, per in enumerate(periods):
                v = p["mois"].get(per["key"])
                cell = wd.cell(row=rr, column=4 + i,
                               value=round(v, p["decimales"]) if v is not None else None)
                cell.number_format = fmt
            cell = wd.cell(row=rr, column=4 + len(periods),
                           value=round(p["total"], p["decimales"])
                           if p["total"] is not None else None)
            cell.number_format, cell.font = fmt, BOLD
            rr += 1
    wd.freeze_panes = "D4"
    wd.column_dimensions["A"].width = 20
    wd.column_dimensions["B"].width = 46
    wd.column_dimensions["C"].width = 8
    for i in range(len(periods) + 1):
        wd.column_dimensions[get_column_letter(4 + i)].width = 13

    # Méthode et points d'attention : le récapitulatif doit porter les mêmes
    # réserves que les classeurs individuels, sinon les chiffres circulent
    # sans elles.
    wm = wb.create_sheet("Méthode")
    add_title(wm, "Méthode et points d'attention", 2)
    f = ev["fraction"]
    rows = [("", "")]
    if ev["source_recommandee"] == "calculee":
        rows += [
            ("Répartition réseau / solaire", "RECALCULÉE"),
            ("Pourquoi",
             "Les sorties réseau/solaire par zone du Miniserver ont été "
             "jugées incohérentes sur la période : " +
             " ; ".join(ev["audit"]["raisons"]) + ". La répartition a donc été "
             "reconstruite à partir des compteurs de l'immeuble."),
            ("Méthode de reconstruction",
             "Pour chaque heure, part solaire du bâtiment = (production - "
             "réinjection) / (achat au réseau + production - réinjection), "
             "appliquée à la consommation mesurée de chaque appartement. "
             "Réseau + solaire redonne exactement la consommation relevée au "
             "compteur du lot."),
            ("Production du bâtiment sur la période",
             f"{_milliers(f['production'])} kWh produits, dont "
             f"{_milliers(f['autoconsomme'])} kWh consommés dans l'immeuble et "
             f"{_milliers(f['injection'])} kWh réinjectés. Achat au réseau : "
             f"{_milliers(f['import_reseau'])} kWh."),
        ]
    else:
        rows += [("Répartition réseau / solaire",
                  "séries du Miniserver (conservées telles quelles)")]
    b = ev["bilan"]
    if b.get("ecart_pct") is not None:
        rows.append(("Bouclage du bilan du bâtiment",
                     f"{b['conso_batiment']:.0f} kWh déduits des compteurs de "
                     f"l'immeuble contre {b['conso_zones']:.0f} kWh aux compteurs "
                     f"de zone, soit {b['ecart_pct']:+.1f} %."))
    rows.append(("Bornes des mois",
                 "Minuit, heure locale suisse (Europe/Zurich), changements "
                 "d'heure inclus."))
    if any(p["role"] == "buanderie" for res in resumes for p in res["postes"]):
        rows.append(("Buanderie",
                     "CEnergieAppXX (kWh) et CNrMachineAppXX (nombre de cycles), "
                     "attribués au lot via le badge NFC. Mise en service Loxone "
                     "le 17.07.2026 -- avant cette date les kWh incrémentaient "
                     "tous les lots (~37 kWh) avec 0 cycle, donc le décompte "
                     "ne part que du 17.07.2026. Un lot à 0 cycle a 0 kWh. "
                     "Les compteurs physiques M1-M6 ne sont pas dans ce "
                     "tableau (pas rattachés à un appartement). Rien n'est "
                     "déduit des Communs."))

    # Un compteur laissé de côté doit se voir ici, sinon son énergie
    # disparaît du décompte sans que personne ne s'en aperçoive.
    non_factures = [(res["label"], p) for res in resumes for p in res["postes"]
                    if p["famille"] == CONTROLE_FAMILY and p["total"] is not None]
    if non_factures:
        rows += [("", ""), ("Compteurs présents mais NON facturés", "")]
        for lab, p in non_factures:
            rows.append((f"{lab} -- {p['titre'].replace(' (compteur de contrôle)', '')}",
                         f"{p['total']:.{p['decimales']}f} {p['unite']} sur la période. "
                         f"Ce compteur n'entre dans aucune colonne du décompte : "
                         f"soit il double un poste déjà facturé, soit son "
                         f"rattachement reste à trancher."))
    rows.append(("Généré le",
                 dt.datetime.fromtimestamp(now_ts, TZ).strftime("%d.%m.%Y à %H:%M")))
    alertes = [(res["label"], a) for res in resumes for a in res["alertes"]]
    if alertes:
        rows += [("", ""), ("Points d'attention par appartement", "")]
        rows += [(lab, a) for lab, a in alertes]
    r = 3
    for a, bb in rows:
        if a and not bb:
            ws_cell = wm.cell(row=r, column=1, value=a)
            ws_cell.font = Font(bold=True, size=12, color="1F4E78")
        else:
            wm.cell(row=r, column=1, value=a).font = BOLD
            wm.cell(row=r, column=2, value=bb).alignment = WRAP_LEFT
        r += 1
    wm.column_dimensions["A"].width = 42
    wm.column_dimensions["B"].width = 100

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def write_recap_csv(path: Path, resumes: list[dict], site: str,
                    periods: list[dict]) -> None:
    """CSV plat de tout le décompte : une ligne par appartement / poste / mois."""
    def fr(v, dec):
        return "" if v is None else f"{round(v, dec):.{dec}f}".replace(".", ",")

    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["Bâtiment", "Appartement", "Poste", "Unité", "Mois",
                    "Mois (libellé)", "Consommation", "Origine"])
        for res in resumes:
            for p in res["postes"]:
                for per in periods:
                    w.writerow([site, res["label"], p["titre"], p["unite"], per["key"],
                                per["label"], fr(p["mois"].get(per["key"]), p["decimales"]),
                                "répartition calculée" if p["calcule"] else "relevé de compteur"])


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("config", help="chemin du fichier config.yaml (donne le db_path)")
    ap.add_argument("--appartement", default=None,
                    help="code d'appartement tel qu'en base (ex: APP35) -- voir /admin")
    ap.add_argument("--tous", action="store_true",
                    help="traiter toutes les zones du site (un classeur par zone "
                         "+ un récapitulatif consolidé)")
    ap.add_argument("--miniserver", default=None,
                    help="site à exporter -- requis si la config en liste plusieurs")
    ap.add_argument("--from-mois", required=True, metavar="AAAA-MM")
    ap.add_argument("--to-mois", required=True, metavar="AAAA-MM")
    ap.add_argument("--out", default=None,
                    help="chemin du .xlsx de sortie (un seul appartement)")
    ap.add_argument("--csv", default=None,
                    help="chemin du .csv de sortie (un seul appartement)")
    ap.add_argument("--outdir", default=None,
                    help="dossier de sortie en mode --tous (défaut : docs/)")
    ap.add_argument("--recap-seul", action="store_true",
                    help="avec --tous : ne produire que le décompte consolidé, "
                         "sans les classeurs individuels par appartement")
    ap.add_argument("--tva", type=float, default=8.1,
                    help="taux de TVA pré-rempli dans le décompte (défaut : 8.1)")
    args = ap.parse_args()

    if not args.appartement and not args.tous:
        raise SystemExit("Préciser --appartement APPxx, ou --tous pour toutes les zones.")
    if args.appartement and args.tous:
        raise SystemExit("--appartement et --tous s'excluent.")

    cfg = load_config(args.config)
    ms_names = [ms.name for ms in cfg.miniservers]
    if args.miniserver:
        if args.miniserver not in ms_names:
            raise SystemExit(f"Miniserver '{args.miniserver}' introuvable. "
                             f"Sites configurés : {', '.join(ms_names)}")
        ms_name = args.miniserver
    elif len(ms_names) == 1:
        ms_name = ms_names[0]
    else:
        raise SystemExit(f"Plusieurs sites configurés ({', '.join(ms_names)}) -- "
                         f"précise --miniserver NOM.")

    conn = db.get_connection(cfg.db_path)
    now_ts = int(time.time())
    series_site = [s for s in db.list_series(conn) if s["miniserver"] == ms_name]
    if not series_site:
        raise SystemExit(f"Aucune série en base pour {ms_name}.")

    y1, m1 = parse_mois(args.from_mois)
    y2, m2 = parse_mois(args.to_mois)
    periods = construire_periodes(y1, m1, y2, m2)
    if not periods:
        raise SystemExit("Plage de mois vide (--from-mois postérieur à --to-mois ?)")

    # Évalué UNE fois pour le site : la crédibilité des séries réseau/solaire
    # est une propriété de l'installation, pas de l'appartement. En mode
    # --tous, ça évite aussi de recalculer la part solaire horaire 22 fois.
    zones = billing.resolve_zones(series_site)
    batiment = billing.resolve_batiment(series_site)
    ev = repartition.evaluer_site(conn, zones, batiment,
                                  periods[0]["start"], periods[-1]["end"])

    print(f"Site        : {ms_name}")
    print(f"Période     : {periods[0]['label']} -> {periods[-1]['label']} "
          f"({len(periods)} mois)")
    print(f"Répartition réseau/solaire : "
          f"{'RECALCULÉE' if ev['source_recommandee'] == 'calculee' else 'séries du Miniserver'}")
    b = ev["bilan"]
    if b.get("ecart_pct") is not None:
        print(f"  bilan du bâtiment : {b['conso_batiment']:.0f} kWh déduits des compteurs "
              f"de bâtiment contre {b['conso_zones']:.0f} kWh aux compteurs de zone "
              f"({b['ecart_pct']:+.1f} %"
              f"{'' if ev['bilan_exploitable'] else ' -- NON exploitable'})")
    for r_ in ev["audit"]["raisons"]:
        print(f"  audit : {r_}")

    if args.tous:
        codes = sorted({(s["apartment"] or "").upper() for s in series_site
                        if (s["apartment"] or "").strip()},
                       key=billing._zone_sort_key)
        print(f"Zones       : {len(codes)} -- {' '.join(codes)}")
        outdir = Path(args.outdir) if args.outdir else Path("docs")
        resumes = []
        for apt in codes:
            res = traiter_appartement(conn, ms_name, series_site, zones, ev, apt,
                                      periods, now_ts, outdir=outdir,
                                      write_files=not args.recap_seul)
            if res:
                resumes.append(res)
        if not resumes:
            raise SystemExit("Aucune zone exploitable sur ce site.")

        suffix = f"{ms_name} {periods[0]['key']} a {periods[-1]['key']}"
        recap = outdir / f"Decompte {suffix}.xlsx"
        build_recap(resumes, ms_name, periods, ev, now_ts, recap, tva=args.tva)
        recap_csv = recap.with_suffix(".csv")
        write_recap_csv(recap_csv, resumes, ms_name, periods)
        if not args.recap_seul:
            print(f"\n{len(resumes)} classeur(s) individuel(s) dans {outdir}/")
        print(f"\nDécompte      : {recap}")
        print(f"CSV consolidé : {recap_csv}")
        return

    apt = args.appartement.strip().upper()
    if not any((s["apartment"] or "").upper() == apt for s in series_site):
        connus = sorted({(s["apartment"] or "?") for s in series_site})
        raise SystemExit(f"Aucune série pour l'appartement '{apt}' sur {ms_name}.\n"
                         f"Appartements connus sur ce site : {', '.join(connus)}")
    res = traiter_appartement(conn, ms_name, series_site, zones, ev, apt, periods,
                              now_ts, out=args.out, csv_path=args.csv)
    if not res:
        raise SystemExit(f"Aucun compteur cumulatif ('total') trouvé pour {apt}.")
    print(f"\nClasseur Excel : {res['xlsx']}")
    print(f"Fichier CSV    : {res['csv']}")


if __name__ == "__main__":
    main()
